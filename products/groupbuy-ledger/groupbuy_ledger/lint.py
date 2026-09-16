"""장부에 쓰인 수식을 검사한다.

두 가지를 본다.

1. **엑셀 2010 에서 없는 함수를 썼는가.** XLOOKUP·FILTER 같은 신형 함수는
   구형 엑셀에서 `#NAME?` 로 깨진다. 고객이 어떤 버전을 쓰는지 모르므로 아예 쓰지 않는다.
2. **수식 글자 안에 오류값이 박혀 있는가.** `#REF!` 가 수식 문자열에 남아 있으면
   시트를 지우다 참조가 끊어진 것이다.

계산해 봐야 아는 오류는 `recalc.py` 가 잡는다. 여기서는 열어 보지 않고 글자만 본다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

__all__ = ["ALLOWED_FUNCTIONS", "BANNED_FUNCTIONS", "LintResult", "scan_formulas"]

#: 엑셀 2010 에 있는 함수만 쓴다. 여기 없는 함수를 쓰려면 2010 에 있는지 확인하고 추가하라.
ALLOWED_FUNCTIONS = frozenset({
    "SUM", "SUMIFS", "SUMIF", "COUNT", "COUNTA", "COUNTIFS", "COUNTIF",
    "INDEX", "MATCH", "IFERROR", "IF", "AND", "OR", "NOT",
    "ROUND", "ROUNDUP", "ROUNDDOWN", "INT", "ABS", "MAX", "MIN", "AVERAGE",
    "RANK", "TEXT", "LEFT", "RIGHT", "MID", "LEN", "TRIM", "CONCATENATE",
    "DATE", "TODAY", "YEAR", "MONTH", "DAY", "EOMONTH", "ROW", "COLUMN",
    "ISBLANK", "ISNUMBER", "ISTEXT", "VALUE", "SUBTOTAL",
})

#: 엑셀 2010 에 없어서 `#NAME?` 로 깨지는 함수. 하나라도 있으면 장부가 잘못 만들어진 것이다.
BANNED_FUNCTIONS = frozenset({
    "XLOOKUP", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY",
    "LET", "LAMBDA", "TEXTJOIN", "TEXTSPLIT", "IFS", "SWITCH", "XMATCH",
    "MAXIFS", "MINIFS", "CONCAT", "ARRAYTOTEXT", "TAKE", "DROP", "VSTACK",
})

#: 수식 안에 이런 글자가 있으면 참조가 이미 깨진 것이다.
ERROR_LITERALS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#NULL!", "#NUM!")

_FUNCTION_CALL = re.compile(r"\b([A-Z][A-Z0-9_.]*)\s*\(")


@dataclass
class LintResult:
    functions: set[str] = field(default_factory=set)
    banned: list[tuple[str, str, str]] = field(default_factory=list)
    unknown: list[tuple[str, str, str]] = field(default_factory=list)
    broken_refs: list[tuple[str, str, str]] = field(default_factory=list)
    formula_count: int = 0

    @property
    def ok(self) -> bool:
        return not (self.banned or self.unknown or self.broken_refs)

    def report(self) -> str:
        if self.ok:
            return (f"수식 {self.formula_count}개 · 쓰인 함수 "
                    f"{', '.join(sorted(self.functions))} — 모두 엑셀 2010 호환")
        lines = []
        for label, rows in (("엑셀 2010 에 없는 함수", self.banned),
                            ("목록에 없는 함수", self.unknown),
                            ("수식에 박힌 오류값", self.broken_refs)):
            if rows:
                lines.append(f"{label} {len(rows)}건")
                for sheet, cell, detail in rows[:10]:
                    lines.append(f"  {sheet}!{cell} → {detail}")
        return "\n".join(lines)


def scan_formulas(path: Path) -> LintResult:
    """통합 문서의 모든 수식을 훑는다."""
    workbook = load_workbook(Path(path), data_only=False)
    result = LintResult()

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                result.formula_count += 1

                for literal in ERROR_LITERALS:
                    if literal in value:
                        result.broken_refs.append(
                            (sheet.title, cell.coordinate, literal))

                # 문자열 안의 글자는 함수 이름으로 세지 않는다 ("취소" 같은 조건값)
                without_text = re.sub(r'"[^"]*"', '""', value)
                for name in _FUNCTION_CALL.findall(without_text):
                    result.functions.add(name)
                    if name in BANNED_FUNCTIONS:
                        result.banned.append((sheet.title, cell.coordinate, name))
                    elif name not in ALLOWED_FUNCTIONS:
                        result.unknown.append((sheet.title, cell.coordinate, name))

    return result
