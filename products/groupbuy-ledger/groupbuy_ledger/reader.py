"""만들어 둔 장부에서 데이터를 도로 읽는다.

`import-orders` 는 기존 파일을 열어 고치지 않고, 여기서 데이터를 뽑아
새 주문을 합친 뒤 통째로 다시 만든다. openpyxl 로 읽고 다시 쓰면
차트가 사라지기 때문이다.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from groupbuy_ledger import layout as L
from groupbuy_ledger.schema import Ledger, Order, Product, Settings

__all__ = ["read_ledger", "LedgerFileError"]


class LedgerFileError(RuntimeError):
    """장부 파일이 이 프로그램이 만든 모양이 아닐 때."""


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    return None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value, default: float = 0) -> float:
    if value is None or isinstance(value, str) and not value.strip():
        return default
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace("원", "").replace("%", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return default
    if isinstance(value, (int, float)):
        return float(value)
    return default


def read_ledger(path: Path) -> Ledger:
    """장부 파일에서 설정·상품·주문을 읽어 모델로 되돌린다."""
    path = Path(path)
    if not path.is_file():
        raise LedgerFileError(f"장부 파일이 없습니다: {path}")

    workbook = load_workbook(path, data_only=False)
    missing = [name for name in L.SHEET_ORDER if name not in workbook.sheetnames]
    if missing:
        raise LedgerFileError(
            f"이 프로그램이 만든 장부가 아닌 것 같습니다. 없는 시트: {', '.join(missing)}"
        )

    config = workbook[L.SHEET_SETTINGS]

    def setting(key: str):
        return config[f"B{L.SET[key]}"].value

    start, end = _as_date(setting("start")), _as_date(setting("end"))
    if start is None or end is None:
        raise LedgerFileError("설정 시트의 시작일·종료일을 날짜로 읽지 못했습니다")

    settings = Settings(
        name=_text(setting("name")) or "이름 없는 공구",
        round_label=_text(setting("round")),
        start=start, end=end,
        default_cost=int(_number(setting("cost"))),
        default_price=int(_number(setting("price"), 1)) or 1,
        shipping_fee=int(_number(setting("shipping_fee"))),
        payment_fee_rate=_number(setting("payment_fee")),
        platform_fee_rate=_number(setting("platform_fee")),
        vat_included=_text(setting("vat_on")) == "예",
        vat_rate=_number(setting("vat_rate"), 0.1),
    )

    products: list[Product] = []
    sheet = workbook[L.SHEET_PRODUCTS]
    for row in sheet.iter_rows(min_row=L.PRODUCT_FIRST_ROW, max_col=5, values_only=True):
        code = _text(row[0])
        if not code:
            continue
        products.append(Product(
            code=code, option=_text(row[1]),
            cost=int(_number(row[2], settings.default_cost)),
            price=int(_number(row[3], settings.default_price)) or settings.default_price,
            initial_stock=int(_number(row[4])),
        ))
    if not products:
        raise LedgerFileError("상품 시트가 비어 있습니다")

    orders: list[Order] = []
    sheet = workbook[L.SHEET_ORDERS]
    for row in sheet.iter_rows(min_row=L.ORDER_FIRST_ROW, max_col=13, values_only=True):
        order_no = _text(row[1])
        order_date = _as_date(row[0])
        if not order_no or order_date is None:
            continue
        orders.append(Order(
            order_date=order_date, order_no=order_no,
            customer=_text(row[2]), phone=_text(row[3]),
            code=_text(row[4]) or "미상", option=_text(row[5]),
            qty=int(_number(row[6], 1)) or 1,
            amount=int(_number(row[7])),
            method=_text(row[8]) or "카드",
            status=_text(row[9]) if _text(row[9]) in
            ("대기", "발송", "완료", "취소", "환불") else "대기",
            invoice=_text(row[10]), memo=_text(row[11]), address=_text(row[12]),
        ))

    return Ledger(settings=settings, products=products, orders=orders)
