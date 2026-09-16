"""만들어진 엑셀을 **엑셀이 아닌 다른 프로그램**으로 다시 계산한다.

왜 필요한가. openpyxl 은 수식을 글자로 써 넣을 뿐 계산하지 않는다.
그래서 "수식이 맞게 들어갔다"와 "열어 보면 맞는 값이 나온다"는 다른 문제다.
`calc.py` 로 파이썬에서 따로 계산한 값과, 이 모듈이 수식을 실제로 돌려서 얻은 값이
같은지 테스트가 맞춰 본다. 두 계산이 서로를 검산하는 셈이다.

엔진은 두 가지를 이 순서로 시도한다.

1. **LibreOffice** — 실제 스프레드시트 프로그램. 있으면 이쪽이 가장 믿을 만하다.
2. **formulas** — 엑셀 수식을 해석해 계산하는 파이썬 라이브러리.

둘 다 없으면 `RecalcUnavailable` 을 낸다. 조용히 넘어가지 않는다.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

__all__ = ["recalculate", "RecalcResult", "RecalcUnavailable", "ERROR_VALUES"]

#: 엑셀이 수식을 못 풀었을 때 칸에 남기는 값. 하나라도 나오면 장부가 깨진 것이다.
ERROR_VALUES = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")

_CELL_KEY = re.compile(r"^'\[[^\]]*\](?P<sheet>[^']+)'!(?P<cell>\$?[A-Z]+\$?\d+)$")


class RecalcUnavailable(RuntimeError):
    """재계산할 수단이 이 컴퓨터에 없을 때."""


@dataclass
class RecalcResult:
    engine: str
    values: dict[tuple[str, str], object] = field(default_factory=dict)

    def get(self, sheet: str, cell: str, default=None):
        return self.values.get((sheet, cell.replace("$", "")), default)

    def errors(self) -> list[tuple[str, str, str]]:
        """(시트, 셀, 오류값) 목록. 완료 기준의 '수식 오류 0건' 을 이걸로 확인한다."""
        found = []
        for (sheet, cell), value in sorted(self.values.items()):
            text = str(value).strip()
            if text in ERROR_VALUES:
                found.append((sheet, cell, text))
        return found


def _libreoffice_binary() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return None


def _recalc_libreoffice(path: Path) -> RecalcResult | None:
    """LibreOffice 로 열었다 저장해 계산값을 캐시에 남긴다."""
    binary = _libreoffice_binary()
    if binary is None:
        return None

    with tempfile.TemporaryDirectory() as workspace:
        profile = Path(workspace) / "profile"
        outdir = Path(workspace) / "out"
        try:
            subprocess.run(
                [binary, "--headless", "--norestore", "--nolockcheck",
                 f"-env:UserInstallation=file://{profile}",
                 "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                check=True, capture_output=True, timeout=300,
            )
        except (subprocess.SubprocessError, OSError):
            return None

        converted = outdir / path.name
        if not converted.is_file():
            return None

        from openpyxl import load_workbook

        workbook = load_workbook(converted, data_only=True)
        result = RecalcResult(engine="libreoffice")
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        result.values[(sheet.title, cell.coordinate)] = cell.value
        return result


def _unwrap(value):
    """formulas 가 돌려주는 Ranges 에서 값 하나를 꺼낸다."""
    raw = getattr(value, "value", value)
    try:
        flat = raw.tolist()
    except AttributeError:
        return raw
    while isinstance(flat, list):
        if not flat:
            return None
        flat = flat[0]
    return flat


def _recalc_formulas(path: Path) -> RecalcResult | None:
    try:
        import formulas
    except ImportError:
        return None

    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()

    result = RecalcResult(engine="formulas")
    for key, value in solution.items():
        found = _CELL_KEY.match(key)
        if not found:
            continue
        sheet = found.group("sheet").strip("'").upper()
        cell = found.group("cell").replace("$", "")
        result.values[(sheet, cell)] = _unwrap(value)

    # formulas 는 시트 이름을 대문자로 돌려준다. 원래 이름으로 맞춰 둔다.
    from openpyxl import load_workbook

    names = {name.upper(): name for name in load_workbook(path).sheetnames}
    result.values = {
        (names.get(sheet, sheet), cell): value
        for (sheet, cell), value in result.values.items()
    }
    return result


def recalculate(path: Path, prefer: str = "") -> RecalcResult:
    """엑셀을 다시 계산한다. 쓸 수 있는 엔진 중 앞선 것을 쓴다."""
    path = Path(path)
    if not path.is_file():
        raise RecalcUnavailable(f"파일이 없습니다: {path}")

    engines = [("libreoffice", _recalc_libreoffice), ("formulas", _recalc_formulas)]
    if prefer:
        engines.sort(key=lambda pair: pair[0] != prefer)

    tried = []
    for name, runner in engines:
        tried.append(name)
        try:
            result = runner(path)
        except Exception:                                 # noqa: BLE001
            result = None                                  # 엔진이 죽으면 다음 것으로
        if result is not None and result.values:
            return result

    raise RecalcUnavailable(
        "엑셀을 다시 계산할 수단이 없습니다 (시도: " + ", ".join(tried) + ").\n"
        "  `pip install formulas` 로 설치하거나 LibreOffice 를 설치하세요."
    )
