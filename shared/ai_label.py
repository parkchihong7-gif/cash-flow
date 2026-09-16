"""AI 생성물 표시 유틸 — 인공지능기본법 제31조(인공지능 생성물 표시) 대응.

이 프로젝트는 AI 프로그램을 **판매**하므로 '인공지능사업자'로서 표시 의무 대상이다
(2026-01-22 시행, 위반 시 시정명령 후 최대 3,000만 원 과태료). 따라서 모든 산출물에
표시를 기본 on 으로 붙인다 (CLAUDE.md §3-5, §7).

- :func:`add_text_label` — 텍스트 산출물 끝에 고지 한 줄을 붙인다.
- :func:`add_metadata` — docx/pptx/xlsx 의 core properties(comments/description)에
  기계 판독용 표시를 기록한다.

두 함수 모두 멱등이다. 이미 표시가 있으면 중복해서 붙이지 않는다.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from shared.config import DEFAULT_TOOL_NAME, PRODUCTS_DIR

__all__ = ["add_text_label", "add_metadata", "label_text_for", "LABELS", "METADATA_PREFIX"]

#: 언어별 고지 문구.
LABELS = {
    "ko": "※ 이 콘텐츠는 생성형 AI의 도움을 받아 제작되었습니다.",
    "en": "* This content was created with the help of generative AI.",
}

#: 메타데이터 표시의 고정 접두사. 중복 기록 판별에도 쓴다.
METADATA_PREFIX = "AI-generated: true;"

_SUPPORTED_SUFFIXES = {".docx", ".pptx", ".xlsx"}


def label_text_for(lang: str = "ko") -> str:
    """언어 코드에 맞는 고지 문구. 모르는 코드는 한국어로 떨어진다."""
    return LABELS.get(lang, LABELS["ko"])


def add_text_label(text: str, lang: str = "ko") -> str:
    """문서 끝에 AI 생성물 고지 한 줄을 붙인다.

    이미 같은 문구가 들어 있으면 그대로 돌려준다 (멱등).
    """
    label = label_text_for(lang)
    if label in text:
        return text
    body = text.rstrip()
    if not body:
        return label
    return f"{body}\n\n{label}"


def _infer_tool(path: Path) -> str:
    """경로가 products/<이름>/... 아래면 그 상품명을 쓴다."""
    try:
        relative = path.resolve().relative_to(Path(PRODUCTS_DIR).resolve())
    except ValueError:
        return DEFAULT_TOOL_NAME
    return relative.parts[0] if relative.parts else DEFAULT_TOOL_NAME


def _existing_marker(value: str) -> str | None:
    """core properties 값에서 이미 기록된 표시 줄을 찾는다."""
    for line in value.splitlines():
        if line.strip().startswith(METADATA_PREFIX):
            return line.strip()
    return None


def _load(path: Path):
    """확장자에 맞는 (문서 객체, core properties, 필드명) 을 돌려준다."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        return document, document.core_properties, "comments"
    if suffix == ".pptx":
        import pptx

        presentation = pptx.Presentation(str(path))
        return presentation, presentation.core_properties, "comments"
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(str(path))
        # openpyxl 의 description 도 docx/pptx 의 comments 와 같은 dc:description 필드다.
        return workbook, workbook.properties, "description"
    raise ValueError(
        f"지원하지 않는 형식입니다: {path.suffix} (지원: {', '.join(sorted(_SUPPORTED_SUFFIXES))})"
    )


def add_metadata(
    path: str | Path,
    tool: str | None = None,
    date: str | None = None,
) -> str:
    """docx/pptx/xlsx 파일에 AI 생성물 표시를 메타데이터로 기록한다.

    core properties 의 comments(= dc:description)에
    ``AI-generated: true; tool: <상품명>; date: <ISO>`` 를 남긴다.

    Args:
        path: 대상 파일. 이미 존재해야 한다.
        tool: 상품명. 생략하면 경로의 `products/<이름>` 에서 추론한다.
        date: ISO 8601 날짜/시각. 생략하면 현재 UTC 시각.

    Returns:
        실제로 기록된 표시 문자열.

    Raises:
        FileNotFoundError: 파일이 없을 때.
        ValueError: 지원하지 않는 확장자일 때.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"파일이 없습니다: {path}")

    tool = tool or _infer_tool(path)
    date = date or datetime.now(timezone.utc).isoformat()
    marker = f"{METADATA_PREFIX} tool: {tool}; date: {date}"

    document, properties, field = _load(path)
    existing = getattr(properties, field) or ""

    already = _existing_marker(existing)
    if already:
        return already  # 이미 표시됨 — 덮어쓰지 않는다

    # python-docx/pptx 는 comments 에 기본 문구를 넣어두므로 지우지 않고 뒤에 덧붙인다.
    setattr(properties, field, f"{existing}\n{marker}".strip() if existing else marker)
    document.save(str(path))
    return marker
