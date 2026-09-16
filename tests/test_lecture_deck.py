"""products/lecture-deck 테스트.

가짜 LLM(`sample_content.fake_ask`)을 주입해 실제 Claude 호출 없이
커리큘럼 → 슬라이드 → pptx/docx/json 전 과정을 검증한다.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import docx
import pytest
from pptx import Presentation
from pptx.util import Emu, Inches
from pydantic import ValidationError

from conftest import load_product_cli
from lecture_deck import docx_builder, pptx_builder
from lecture_deck.generator import (
    BASE_DIR, DeckGenerator, Slide, assemble_deck, pad_note,
)
from lecture_deck.sample_content import fake_ask
from lecture_deck.schema import (
    DeckInput, MAX_BODY_LINES, MAX_LINE_CHARS, MAX_MODULES, MIN_MODULES,
    NOTE_MAX_CHARS, NOTE_MIN_CHARS, STYLES, fit_body, load_input, wrap_line,
)
from shared import banned_phrases

deck_cli = load_product_cli("lecture-deck")

DECK_DIR = Path(BASE_DIR)
EXAMPLE_YAML = DECK_DIR / "deck_input.yaml"


# ------------------------------------------------------------------ 픽스처
def _input(**overrides) -> DeckInput:
    payload = {
        "course_title": "테스트 강의",
        "audience": "테스트 수강생",
        "total_minutes": 180,
        "style": "minimal",
        "modules": [],
    }
    payload.update(overrides)
    return DeckInput(**payload)


@pytest.fixture
def deck():
    """전 모듈 슬라이드까지 조립된 덱."""
    data = _input()
    generator = DeckGenerator(data, model="test", ask_fn=fake_ask)
    curriculum = generator.build_curriculum()

    module_slides, previous = {}, ""
    for index in range(1, len(curriculum.modules) + 1):
        result = generator.write_module(curriculum, index, previous)
        assert result.ok, f"{index}번 모듈 실패"
        module_slides[index] = result.slides
        summary = next((s for s in result.slides if s.kind == "summary"), None)
        previous = " / ".join(summary.body) if summary else ""

    return assemble_deck(data, curriculum, module_slides, "test")


@pytest.fixture
def built(tmp_path, deck):
    pptx_path = pptx_builder.build_pptx(deck, tmp_path / "deck.pptx")
    docx_path = docx_builder.build_worksheet(deck, tmp_path / "worksheet.docx")
    deck_cli.write_curriculum(deck, tmp_path)
    json_path = deck_cli.write_outline_json(deck, tmp_path)
    return {"dir": tmp_path, "pptx": pptx_path, "docx": docx_path,
            "json": json_path, "deck": deck}


# --------------------------------------------------------------- 입력 검증
def test_example_yaml_loads():
    data = load_input(EXAMPLE_YAML)
    assert data.style in STYLES
    assert data.modules == [], "예시는 모듈을 비워 LLM 제안을 받도록 둔다"


@pytest.mark.parametrize(
    ("overrides", "reason"),
    [
        ({"total_minutes": 20}, "너무 짧은 강의"),
        ({"total_minutes": 1000}, "너무 긴 강의"),
        ({"style": "fancy"}, "없는 스타일"),
        ({"course_title": "  "}, "빈 제목"),
        ({"modules": [{"title": f"m{i}"} for i in range(4)]}, "모듈 5개 미만"),
        ({"modules": [{"title": f"m{i}"} for i in range(9)]}, "모듈 8개 초과"),
    ],
)
def test_invalid_input_rejected(overrides, reason):
    with pytest.raises(ValidationError):
        _input(**overrides)


def test_module_concept_count_enforced():
    with pytest.raises(ValidationError, match="핵심 개념"):
        _input(modules=[{"title": "m", "concepts": ["가", "나"]}] * 5)


def test_unknown_key_rejected():
    with pytest.raises(ValidationError):
        _input(강의료=100000)


def test_duration_text():
    assert _input(total_minutes=180).duration_text == "3시간"
    assert _input(total_minutes=90).duration_text == "1시간 30분"
    assert _input(total_minutes=45).duration_text == "45분"


# ------------------------------------------------------- 텍스트 맞춤 (핵심)
@pytest.mark.parametrize(
    "text",
    [
        "이건 스무 글자를 훌쩍 넘기는 아주 긴 본문 한 줄입니다",
        "짧은 줄",
        "띄어쓰기없이아주긴한단어가들어있는경우도처리해야합니다정말로",
        "",
    ],
)
def test_wrap_line_never_exceeds_limit(text):
    assert all(len(line) <= MAX_LINE_CHARS for line in wrap_line(text))


def test_wrap_line_keeps_all_characters():
    text = "이건 스무 글자를 훌쩍 넘기는 긴 줄입니다"
    assert "".join(wrap_line(text)).replace(" ", "") == text.replace(" ", "")


def test_fit_body_splits_into_slides():
    lines = [f"줄 {i}" for i in range(1, 16)]
    chunks = fit_body(lines)
    assert len(chunks) == 3
    assert all(len(chunk) <= MAX_BODY_LINES for chunk in chunks)


def test_fit_body_handles_empty():
    assert fit_body([]) == []
    assert fit_body(["", "  "]) == []


# ------------------------------------------------- 완료 기준: 6줄·20자
def test_every_slide_obeys_body_rules(deck):
    """완료 기준: 모든 슬라이드 본문이 6줄·20자 규칙을 통과한다."""
    violations = [
        f"{index}번({slide.title}): {len(slide.body)}줄, "
        f"최대 {max((len(l) for l in slide.body), default=0)}자"
        for index, slide in enumerate(deck.slides, start=1)
        if slide.overflows
    ]
    assert violations == [], f"규칙 위반 슬라이드: {violations}"


def test_overflowing_slide_is_split_automatically():
    """LLM 이 규칙을 어겨도 산출물에는 위반이 없어야 한다."""
    data = _input()
    generator = DeckGenerator(data, model="test", ask_fn=fake_ask)
    curriculum = generator.build_curriculum()

    fat = Slide(
        kind="concept", title="넘치는 슬라이드",
        body=["이건 스무 글자를 훌쩍 넘기는 아주 긴 줄입니다"] * 5,
        note="원래 노트",
    )
    result = assemble_deck(data, curriculum, {1: [fat]}, "test")

    assert not any(slide.overflows for slide in result.slides)
    assert result.split_count > 0
    parts = [s for s in result.slides if s.title.startswith("넘치는 슬라이드")]
    assert len(parts) > 1
    assert "(1/" in parts[0].title


# --------------------------------------- 완료 기준: 발표자 노트 빈 곳 0장
def test_no_slide_has_an_empty_note(deck):
    """완료 기준: 발표자 노트가 비어 있는 슬라이드 0장."""
    empty = [
        f"{index}번({slide.title})"
        for index, slide in enumerate(deck.slides, start=1)
        if not slide.note.strip()
    ]
    assert empty == [], f"노트가 빈 슬라이드: {empty}"


def test_every_note_is_within_range(deck):
    for index, slide in enumerate(deck.slides, start=1):
        assert NOTE_MIN_CHARS <= len(slide.note) <= NOTE_MAX_CHARS, (
            f"{index}번({slide.title}) 노트 {len(slide.note)}자"
        )


@pytest.mark.parametrize("note", ["", "   ", "너무 짧은 노트"])
def test_pad_note_fills_short_notes(note):
    slide = Slide(kind="concept", title="제목", body=["가", "나"])
    result = pad_note(note, slide, "강의명")
    assert NOTE_MIN_CHARS <= len(result) <= NOTE_MAX_CHARS
    assert "강사가 채울 곳" in result


def test_pad_note_keeps_good_notes():
    good = "가" * 250
    slide = Slide(kind="concept", title="제목")
    assert pad_note(good, slide, "강의명") == good


def test_pad_note_truncates_long_notes():
    slide = Slide(kind="concept", title="제목")
    assert len(pad_note("가" * 900, slide, "강의명")) == NOTE_MAX_CHARS


# ---------------------------------------------------------- 덱 구조
def test_deck_has_the_required_fixed_slides(deck):
    kinds = [slide.kind for slide in deck.slides]
    assert kinds[0] == "cover"
    assert kinds[1] == "instructor"
    assert kinds[2] == "agenda"
    assert kinds[-2] == "closing"
    assert kinds[-1] == "ai_notice"


def test_each_module_has_cover_concepts_practice_summary(deck):
    for number in range(1, len(deck.curriculum.modules) + 1):
        kinds = [s.kind for s in deck.slides if s.module_number == number]
        assert kinds[0] == "module", f"{number}번 모듈 표지 없음"
        assert kinds.count("practice") == 1, f"{number}번 실습 슬라이드"
        assert kinds.count("summary") == 1, f"{number}번 요약 슬라이드"
        assert kinds.count("concept") >= 3, f"{number}번 개념 슬라이드"


def test_deck_is_about_forty_slides(deck):
    """완료 기준: 예시 입력으로 40장 내외."""
    assert 35 <= deck.slide_count <= 50, f"{deck.slide_count}장"


def test_agenda_lists_every_module(deck):
    agenda = next(s for s in deck.slides if s.kind == "agenda")
    assert len(agenda.body) >= len(deck.curriculum.modules)


def test_instructor_slide_leaves_placeholders(deck):
    slide = next(s for s in deck.slides if s.kind == "instructor")
    assert any("[" in line for line in slide.body)


def test_curriculum_module_count_in_range(deck):
    assert MIN_MODULES <= len(deck.curriculum.modules) <= MAX_MODULES


def test_previous_module_summary_is_passed_on():
    data = _input()
    generator = DeckGenerator(data, model="test", ask_fn=fake_ask)
    curriculum = generator.build_curriculum()

    seen: list[str] = []

    def spy(system, user, model="test", json_mode=False, **kw):
        seen.append(user)
        return fake_ask(system, user, model, json_mode)

    spy_generator = DeckGenerator(data, model="test", ask_fn=spy)
    spy_generator.write_module(curriculum, 1, "")
    spy_generator.write_module(curriculum, 2, "앞 모듈 요약입니다")

    assert "[앞 모듈 요약]" not in seen[0]
    assert "앞 모듈 요약입니다" in seen[1]


def test_module_failure_is_retried_alone():
    from shared.llm import LLMError

    data = _input()
    curriculum = DeckGenerator(data, model="test", ask_fn=fake_ask).build_curriculum()
    state = {"calls": 0}

    def flaky(system, user, model="test", json_mode=False, **kw):
        state["calls"] += 1
        if state["calls"] == 1:
            raise LLMError("일시적인 오류")
        return fake_ask(system, user, model, json_mode)

    result = DeckGenerator(data, model="test", ask_fn=flaky).write_module(
        curriculum, 1, ""
    )
    assert result.ok and result.attempts == 2


def test_banned_phrase_triggers_rewrite():
    data = _input()
    curriculum = DeckGenerator(data, model="test", ask_fn=fake_ask).build_curriculum()
    state = {"calls": 0}

    def dirty_first(system, user, model="test", json_mode=False, **kw):
        state["calls"] += 1
        payload = fake_ask(system, user, model, json_mode)
        if state["calls"] == 1:
            payload["slides"][0]["title"] = "수익 보장 비법"
        return payload

    result = DeckGenerator(data, model="test", ask_fn=dirty_first).write_module(
        curriculum, 1, ""
    )
    assert result.attempts == 2
    assert result.banned == []


# ------------------------------------------------------------------ pptx
def test_pptx_is_created_and_reopens(built):
    presentation = Presentation(str(built["pptx"]))
    assert len(presentation.slides) == built["deck"].slide_count


def test_pptx_is_16_by_9(built):
    presentation = Presentation(str(built["pptx"]))
    ratio = presentation.slide_width / presentation.slide_height
    assert abs(ratio - 16 / 9) < 0.01
    assert round(Emu(presentation.slide_width).inches, 2) == 13.33


def test_pptx_notes_are_never_empty(built):
    """완료 기준: 발표자 노트가 비어 있는 슬라이드 0장 (pptx 실물 기준)."""
    presentation = Presentation(str(built["pptx"]))
    empty = [
        index for index, slide in enumerate(presentation.slides, start=1)
        if not slide.notes_slide.notes_text_frame.text.strip()
    ]
    assert empty == []


def test_pptx_shapes_stay_inside_the_canvas(built):
    """렌더링을 못 하는 환경이라 좌표로 오버플로를 확인한다."""
    presentation = Presentation(str(built["pptx"]))
    width, height = presentation.slide_width, presentation.slide_height
    outside = []
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > width
                    or shape.top + shape.height > height):
                outside.append(f"{index}번 {shape.shape_id}")
    assert outside == []


def test_pptx_respects_half_inch_margins(built):
    presentation = Presentation(str(built["pptx"]))
    margin = Inches(0.5)
    too_close = [
        f"{index}번"
        for index, slide in enumerate(presentation.slides, start=1)
        for shape in slide.shapes
        if shape.left is not None and shape.left < margin
    ]
    assert too_close == []


def test_pptx_uses_no_external_images(built):
    """사양: python-pptx 기본 도형만. 외부 이미지에 의존하지 않는다."""
    with zipfile.ZipFile(built["pptx"]) as archive:
        media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
    assert media == []


def test_pptx_sets_korean_font(built):
    with zipfile.ZipFile(built["pptx"]) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "맑은 고딕" in slide_xml
    assert "typeface" in slide_xml


def test_pptx_carries_slide_text(built):
    presentation = Presentation(str(built["pptx"]))
    text = "\n".join(
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert built["deck"].data.course_title in text
    assert "AI 활용 고지" in text


def test_pptx_has_no_banned_phrases(built):
    presentation = Presentation(str(built["pptx"]))
    text = "\n".join(
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    notes = "\n".join(
        slide.notes_slide.notes_text_frame.text for slide in presentation.slides
    )
    assert banned_phrases.check(text + notes) == []


@pytest.mark.parametrize("style", STYLES)
def test_every_style_renders(tmp_path, style):
    data = _input(style=style)
    generator = DeckGenerator(data, model="test", ask_fn=fake_ask)
    curriculum = generator.build_curriculum()
    result = generator.write_module(curriculum, 1, "")
    deck = assemble_deck(data, curriculum, {1: result.slides}, "test")

    path = pptx_builder.build_pptx(deck, tmp_path / f"{style}.pptx")
    assert path.is_file() and path.stat().st_size > 10_000


def test_unknown_style_is_rejected():
    with pytest.raises(ValueError, match="모르는 스타일"):
        pptx_builder.load_palette("없는스타일")


@pytest.mark.parametrize("style", STYLES)
def test_palette_has_every_required_color(style):
    palette = pptx_builder.load_palette(style)
    for key in ("dark_bg", "light_bg", "panel_bg", "ink", "ink_on_dark",
                "muted", "muted_on_dark", "accent", "accent_soft", "on_accent"):
        assert key in palette, f"{style}: {key} 없음"
        assert len(palette[key]) == 6, f"{style}.{key} 는 6자리 hex 여야 합니다"


# ------------------------------------------------------------ worksheet
def test_worksheet_is_created(built):
    document = docx.Document(str(built["docx"]))
    assert len(document.paragraphs) > 50


def test_worksheet_has_three_blanks_per_module(built):
    document = docx.Document(str(built["docx"]))
    text = "\n".join(p.text for p in document.paragraphs)
    modules = built["deck"].curriculum.modules
    for index, module in enumerate(modules, start=1):
        assert f"모듈 {index}. {module.get('title', '')}" in text
    assert text.count("＿" * 18) >= len(modules) * 3


def test_worksheet_has_checkboxes(built):
    text = "\n".join(p.text for p in docx.Document(str(built["docx"])).paragraphs)
    assert text.count("☐") >= len(built["deck"].curriculum.modules) * 3


def test_worksheet_records_ai_metadata(built):
    core = docx.Document(str(built["docx"])).core_properties
    assert "AI-generated: true" in (core.comments or "")
    assert "tool: lecture-deck" in core.comments


def test_worksheet_settings_pass_schema(built):
    with zipfile.ZipFile(built["docx"]) as archive:
        settings = archive.read("word/settings.xml").decode("utf-8")
    assert 'w:val="bestFit"' not in settings


# --------------------------------------------------- curriculum / json
def test_curriculum_md_lists_every_module(built):
    text = (built["dir"] / "curriculum.md").read_text(encoding="utf-8")
    for module in built["deck"].curriculum.modules:
        assert module["title"] in text
        assert module["practice"] in text
        assert module["assessment"] in text
    assert "강의 전 확인" in text
    assert banned_phrases.check(text) == []


def test_curriculum_md_warns_on_time_mismatch(tmp_path, deck):
    deck.curriculum.modules[0]["minutes"] = 999
    deck_cli.write_curriculum(deck, tmp_path)
    text = (tmp_path / "curriculum.md").read_text(encoding="utf-8")
    assert "목표" in text and "다릅니다" in text


def test_outline_json_has_every_slide(built):
    payload = json.loads(built["json"].read_text(encoding="utf-8"))
    assert len(payload["slides"]) == built["deck"].slide_count
    assert payload["limits"]["max_line_chars"] == MAX_LINE_CHARS
    for slide in payload["slides"]:
        assert slide["note"].strip(), f"{slide['index']}번 노트 비어 있음"
        assert len(slide["body"]) <= MAX_BODY_LINES


def test_outline_json_is_reusable(built):
    """다른 도구에서 쓸 수 있게 모듈 정보까지 담는다."""
    payload = json.loads(built["json"].read_text(encoding="utf-8"))
    assert payload["modules"]
    assert payload["course_title"] and payload["style"]


# ------------------------------------------------------------------ 문서
def test_readme_covers_workflow_and_review():
    text = (DECK_DIR / "README.md").read_text(encoding="utf-8")
    for token in ["deck_input.yaml", "deck.pptx", "worksheet.docx", "palettes.json"]:
        assert token in text, token
    assert "검수" in text
    assert banned_phrases.check(text) == []


def test_product_manuals_are_detailed_and_clean():
    for name in ("admin.md", "client.md"):
        path = DECK_DIR / "docs" / name
        assert path.is_file(), name
        text = path.read_text(encoding="utf-8")
        assert len(text) > 4000, f"{name} 이 너무 짧습니다"
        assert banned_phrases.check(text) == []


def test_prompts_state_the_hard_limits():
    slides = (DECK_DIR / "prompts" / "slides.md").read_text(encoding="utf-8")
    assert str(MAX_LINE_CHARS) in slides
    assert str(MAX_BODY_LINES) in slides
    assert str(NOTE_MIN_CHARS) in slides
