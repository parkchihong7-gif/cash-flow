"""products/groupbuy-ledger 테스트.

가장 중요한 것은 **엑셀이 실제로 계산한 값**과 파이썬이 따로 계산한 값이
같은지 보는 것이다(`test_recalculated_profit_matches_python`).
openpyxl 은 수식을 글자로 써 넣을 뿐 계산하지 않으므로,
수식이 들어갔다는 것만으로는 "열어 보면 맞다" 가 증명되지 않는다.

재계산 엔진은 LibreOffice 를 먼저 쓰고, 없으면 `formulas` 라이브러리를 쓴다.
둘 다 없으면 테스트는 **건너뛰지 않고 실패한다.** 완료 기준이 걸린 검증이라
조용히 넘어가면 안 된다.
"""

from __future__ import annotations

import csv
import json
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from conftest import load_product_cli
from groupbuy_ledger import layout as L
from groupbuy_ledger.calc import compute, excel_round
from groupbuy_ledger.importer import (
    ImportError_, import_orders, load_mapping, read_orders_csv, read_products_csv,
)
from groupbuy_ledger.lint import BANNED_FUNCTIONS, scan_formulas
from groupbuy_ledger.reader import LedgerFileError, read_ledger
from groupbuy_ledger.recalc import recalculate
from groupbuy_ledger.report import (
    SHIPPING_HEADER, gather, insight_lines, unknown_numbers,
    write_report, write_shipping_csv,
)
from groupbuy_ledger.sample_content import fake_ask
from groupbuy_ledger.schema import (
    STOCK_WARN, Ledger, Order, Product, Settings, slugify,
)
from groupbuy_ledger.workbook import build_workbook
from shared import banned_phrases

gb_cli = load_product_cli("groupbuy-ledger")

BASE_DIR = Path(gb_cli.BASE_DIR)
SAMPLES = BASE_DIR / "samples"
MAPPING = BASE_DIR / "mapping.yaml"


# ------------------------------------------------------------------ 픽스처
def _settings(**overrides) -> Settings:
    base = dict(name="9월 원두 공구", start=date(2026, 9, 1), end=date(2026, 9, 14),
                default_cost=9000, default_price=15000, shipping_fee=3000,
                payment_fee_rate=0.032, platform_fee_rate=0.0, vat_included=True)
    base.update(overrides)
    return Settings(**base)


@pytest.fixture(scope="session")
def sample_ledger() -> Ledger:
    """샘플 CSV 두 개로 만든 장부. 완료 기준의 new → import 경로와 같다."""
    products = [Product(**row) for row in
                read_products_csv(SAMPLES / "products.csv", 9000, 15000)]
    ledger = Ledger(settings=_settings(), products=products, orders=[])
    merged, _ = import_orders(ledger, SAMPLES / "orders.csv", MAPPING)
    return merged


@pytest.fixture(scope="session")
def built(tmp_path_factory, sample_ledger) -> Path:
    path = tmp_path_factory.mktemp("ledger") / sample_ledger.filename
    return build_workbook(sample_ledger, path)


@pytest.fixture(scope="session")
def recalculated(built):
    """엑셀을 실제로 다시 계산한 결과. 세션에 한 번만 돌린다(느리다)."""
    return recalculate(built)


# ------------------------------------------------------------------ 스키마
def test_margin_rate_is_one_minus_cost_over_price():
    settings = _settings(default_cost=9000, default_price=15000)
    assert settings.margin_rate == pytest.approx(0.4)


def test_settings_rejects_backwards_period():
    with pytest.raises(ValidationError, match="종료일"):
        _settings(start=date(2026, 9, 14), end=date(2026, 9, 1))


def test_settings_rejects_price_below_cost():
    with pytest.raises(ValidationError, match="판매가"):
        _settings(default_cost=20000, default_price=10000)


def test_round_label_defaults_to_the_name():
    assert _settings(name="9월 원두 공구").round_label == "9월-원두-공구"


def test_slugify_drops_characters_that_break_filenames():
    assert "/" not in slugify("9월/공구")
    assert slugify("  ") == "회차"


def test_ledger_rejects_duplicate_products():
    product = Product(code="A", option="기본", cost=1000, price=2000, initial_stock=5)
    with pytest.raises(ValidationError, match="겹칩니다"):
        Ledger(settings=_settings(), products=[product, product])


def test_ledger_rejects_duplicate_order_numbers(sample_ledger):
    order = sample_ledger.orders[0]
    with pytest.raises(ValidationError, match="주문번호가 겹칩니다"):
        Ledger(settings=sample_ledger.settings, products=sample_ledger.products,
               orders=[order, order.model_copy()])


def test_unknown_order_keys_are_reported(sample_ledger):
    stray = sample_ledger.orders[0].model_copy(
        update={"order_no": "STRAY-1", "code": "없는코드"})
    ledger = Ledger(settings=sample_ledger.settings, products=sample_ledger.products,
                    orders=[*sample_ledger.orders, stray])
    assert any("없는코드" in key for key in ledger.unknown_order_keys())


# ------------------------------------------------------------------ 계산
def test_excel_round_rounds_half_away_from_zero():
    """파이썬 기본 round 는 2.5 를 2 로 만든다. 엑셀은 3 으로 만든다."""
    assert excel_round(2.5) == 3
    assert excel_round(0.5) == 1
    assert round(2.5) == 2


def test_void_orders_are_excluded_from_revenue():
    products = [Product(code="A", option="기본", cost=1000, price=2000, initial_stock=10)]
    orders = [
        Order(order_date=date(2026, 9, 1), order_no="1", code="A", option="기본",
              qty=2, amount=4000, status="완료"),
        Order(order_date=date(2026, 9, 1), order_no="2", code="A", option="기본",
              qty=1, amount=2000, status="취소"),
        Order(order_date=date(2026, 9, 1), order_no="3", code="A", option="기본",
              qty=1, amount=2000, status="환불"),
    ]
    totals = compute(Ledger(settings=_settings(default_cost=1000, default_price=2000),
                            products=products, orders=orders))
    assert totals.gross == 8000
    assert totals.void == 4000
    assert totals.net == 4000
    assert totals.products[0].sold == 2, "취소·환불 수량은 판매로 세지 않는다"
    assert totals.products[0].stock == 8, "취소된 주문은 재고를 까지 않는다"


def test_vat_can_be_switched_off():
    ledger = Ledger(settings=_settings(vat_included=False),
                    products=[Product(code="A", option="기본", cost=1000,
                                      price=2000, initial_stock=10)],
                    orders=[Order(order_date=date(2026, 9, 1), order_no="1", code="A",
                                  option="기본", qty=1, amount=2000, status="완료")])
    assert compute(ledger).vat == 0


def test_ranks_are_unique_even_when_profits_tie():
    products = [Product(code=f"P{i}", option="기본", cost=1000, price=2000,
                        initial_stock=10) for i in range(3)]
    orders = [Order(order_date=date(2026, 9, 1), order_no=str(i), code=f"P{i}",
                    option="기본", qty=1, amount=2000, status="완료") for i in range(3)]
    totals = compute(Ledger(settings=_settings(default_cost=1000, default_price=2000),
                            products=products, orders=orders))
    assert sorted(p.rank for p in totals.products) == [1, 2, 3]


def test_empty_ledger_does_not_divide_by_zero():
    totals = compute(Ledger(
        settings=_settings(),
        products=[Product(code="A", option="기본", cost=1, price=2, initial_stock=0)],
        orders=[]))
    assert totals.margin is None and totals.refund_rate is None
    assert totals.avg_order is None and totals.profit == 0


# ------------------------------------------------------------------ 가져오기
def test_sample_csvs_exist_and_load():
    """완료 기준: 샘플 products.csv, orders.csv 포함."""
    products = read_products_csv(SAMPLES / "products.csv", 0, 1)
    assert len(products) >= 5
    with (SAMPLES / "orders.csv").open(encoding="utf-8") as handle:
        assert len(list(csv.DictReader(handle))) >= 20


def test_import_maps_differently_named_columns(sample_ledger):
    """샘플 주문 CSV 는 '타임스탬프'·'구매자명'·'주문상태' 처럼 이름이 다르다."""
    result = read_orders_csv(SAMPLES / "orders.csv", load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert result.matched_columns["주문일"] == "타임스탬프"
    assert result.matched_columns["고객명"] == "구매자명"
    assert result.matched_columns["배송상태"] == "주문상태"
    assert result.matched_columns["메모"] == "배송메시지"
    assert not result.missing_columns


def test_import_parses_money_and_keeps_phone_text(sample_ledger):
    order = sample_ledger.orders[0]
    assert order.amount > 0, "'15,000원' 을 숫자로 읽어야 한다"
    assert order.phone.startswith("010-"), "연락처의 앞자리 0 이 남아야 한다"


def test_import_skips_duplicate_order_numbers(sample_ledger):
    """같은 CSV 를 두 번 넣어도 주문이 늘어나지 않는다."""
    merged, result = import_orders(sample_ledger, SAMPLES / "orders.csv", MAPPING)
    assert result.added == []
    assert len(result.duplicated) == len(sample_ledger.orders)
    assert len(merged.orders) == len(sample_ledger.orders)
    assert result.needs_attention


def test_import_normalizes_shop_status_names(tmp_path, sample_ledger):
    path = tmp_path / "shop.csv"
    path.write_text(
        "결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
        "2026-09-02,X1,COFFEE-01,에티오피아 200g,1,\"15,000\",배송완료\n"
        "2026-09-02,X2,COFFEE-01,에티오피아 200g,1,\"15,000\",상품준비중\n"
        "2026-09-02,X3,COFFEE-01,에티오피아 200g,1,\"15,000\",반품완료\n",
        encoding="utf-8")
    result = read_orders_csv(path, load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert [o.status for o in result.added] == ["완료", "대기", "환불"]


def test_import_flags_unknown_status_instead_of_guessing(tmp_path, sample_ledger):
    path = tmp_path / "odd.csv"
    path.write_text(
        "결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
        "2026-09-02,Y1,COFFEE-01,에티오피아 200g,1,15000,알수없는상태\n",
        encoding="utf-8")
    result = read_orders_csv(path, load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert result.added[0].status == "대기"
    assert "알수없는상태" in result.unknown_status


def test_import_flags_products_missing_from_the_product_sheet(tmp_path, sample_ledger):
    path = tmp_path / "stray.csv"
    path.write_text(
        "결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
        "2026-09-02,Z1,없는상품,특대,1,15000,완료\n", encoding="utf-8")
    result = read_orders_csv(path, load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert any("없는상품" in item for item in result.unknown_products)
    assert result.needs_attention


def test_import_names_the_missing_column_and_shows_the_csv_headers(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("이상한칸,또다른칸\n1,2\n", encoding="utf-8")
    with pytest.raises(ImportError_) as excinfo:
        read_orders_csv(path, load_mapping(MAPPING), known_keys=set(), existing_numbers=set())
    message = str(excinfo.value)
    assert "주문번호" in message
    assert "이상한칸" in message, "CSV 에 실제로 있던 컬럼을 알려줘야 고칠 수 있다"


def test_import_reads_cp949_files(tmp_path, sample_ledger):
    """스마트스토어는 cp949 로 내려주는 경우가 많다."""
    path = tmp_path / "cp949.csv"
    path.write_bytes(
        ("결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
         "2026-09-02,C1,COFFEE-01,에티오피아 200g,1,15000,배송완료\n").encode("cp949"))
    result = read_orders_csv(path, load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert result.added[0].code == "COFFEE-01"


def test_import_skips_rows_with_an_unreadable_date(tmp_path, sample_ledger):
    path = tmp_path / "nodate.csv"
    path.write_text(
        "결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
        "날짜아님,D1,COFFEE-01,에티오피아 200g,1,15000,완료\n", encoding="utf-8")
    result = read_orders_csv(path, load_mapping(MAPPING),
                             known_keys=sample_ledger.product_keys, existing_numbers=set())
    assert result.added == []
    assert result.skipped_rows


# ------------------------------------------------------------------ 통합문서
def test_workbook_has_the_six_sheets_in_order(built):
    assert load_workbook(built).sheetnames == list(L.SHEET_ORDER)


def test_orders_sheet_has_the_twelve_columns_in_the_given_order(built):
    sheet = load_workbook(built)[L.SHEET_ORDERS]
    headers = [sheet[f"{letter}1"].value for letter, _, _ in L.ORDER_COLUMNS]
    assert headers[:12] == ["주문일", "주문번호", "고객명", "연락처", "상품코드", "옵션",
                            "수량", "결제금액", "결제수단", "배송상태", "송장번호", "메모"]


def test_order_status_has_a_dropdown_with_five_choices(built):
    sheet = load_workbook(built)[L.SHEET_ORDERS]
    lists = [dv for dv in sheet.data_validations.dataValidation if dv.type == "list"]
    statuses = [dv for dv in lists if "대기" in str(dv.formula1)]
    assert statuses, "배송상태 드롭다운이 없습니다"
    assert str(statuses[0].formula1).strip('"').split(",") == \
        ["대기", "발송", "완료", "취소", "환불"]


def test_product_code_dropdown_points_at_the_product_sheet(built):
    sheet = load_workbook(built)[L.SHEET_ORDERS]
    formulas = [str(dv.formula1) for dv in sheet.data_validations.dataValidation]
    assert any(L.SHEET_PRODUCTS in f and "$A$2" in f for f in formulas)


def test_void_rows_and_low_stock_have_conditional_formatting(built):
    workbook = load_workbook(built)
    orders = str(workbook[L.SHEET_ORDERS].conditional_formatting._cf_rules)
    assert "취소" in orders and "환불" in orders, "취소·환불 행 서식이 없습니다"

    products = workbook[L.SHEET_PRODUCTS].conditional_formatting
    rules = [rule for rules in products._cf_rules.values() for rule in rules]
    assert rules, "재고 부족 서식이 없습니다"
    assert any(str(STOCK_WARN) in str(rule.formula) for rule in rules)


def test_dashboard_has_both_charts(built):
    with zipfile.ZipFile(built) as archive:
        charts = [n for n in archive.namelist() if n.startswith("xl/charts/chart")]
    assert len(charts) >= 2, "일별 매출 추이와 상품별 이익 순위 차트가 있어야 합니다"


def test_shipping_sheet_uses_the_courier_column_order(built):
    sheet = load_workbook(built)[L.SHEET_SHIPPING]
    headers = [sheet[f"{letter}1"].value for letter, _, _ in L.SHIPPING_COLUMNS]
    assert headers[:5] == ["받는분성명", "받는분전화번호", "받는분주소", "품목명", "수량"]


def test_shipping_sheet_pulls_rows_with_index_match(built):
    sheet = load_workbook(built)[L.SHEET_SHIPPING]
    formula = sheet["A2"].value
    assert formula.startswith("=IFERROR(INDEX(")
    assert "MATCH(" in formula and L.SHEET_ORDERS in formula


def test_ledger_sheet_totals_are_formulas_not_baked_numbers(built):
    sheet = load_workbook(built)[L.SHEET_LEDGER]
    for key, row in L.SUMMARY.items():
        value = sheet[f"B{row}"].value
        assert isinstance(value, str) and value.startswith("="), f"{key} 가 수식이 아닙니다"


def test_product_table_uses_sumifs(built):
    sheet = load_workbook(built)[L.SHEET_LEDGER]
    row = L.PRODUCT_TABLE_FIRST_ROW
    assert "SUMIFS(" in sheet[f"C{row}"].value
    assert "SUMIFS(" in sheet[f"D{row}"].value


def test_settings_margin_is_one_minus_cost_over_price_as_a_formula(built):
    sheet = load_workbook(built)[L.SHEET_SETTINGS]
    assert sheet[f"B{L.SET['margin']}"].value.startswith("=IFERROR(1-")


# ------------------------------------------------- 엑셀 2010 호환·수식 오류
def test_only_excel_2010_functions_are_used(built):
    """XLOOKUP·FILTER 를 쓰면 구형 엑셀에서 #NAME? 로 깨진다."""
    result = scan_formulas(built)
    assert result.banned == [], result.report()
    assert result.unknown == [], result.report()
    assert result.functions & BANNED_FUNCTIONS == set()
    assert result.formula_count > 100


def test_no_error_literals_are_baked_into_formulas(built):
    assert scan_formulas(built).broken_refs == []


def test_no_formula_errors_after_recalculation(recalculated):
    """완료 기준: 수식 오류(#REF!, #NAME?) 0건."""
    errors = recalculated.errors()
    assert errors == [], f"수식 오류 {len(errors)}건: {errors[:10]}"


# ----------------------------------------------- 엑셀 계산값 vs 파이썬 계산값
def test_recalculated_profit_matches_python(recalculated, sample_ledger):
    """완료 기준: 재계산한 정산 시트 순이익이 파이썬 독립 계산과 일치."""
    expected = compute(sample_ledger).profit
    actual = recalculated.get(L.SHEET_LEDGER, f"B{L.SUMMARY['profit']}")
    assert actual is not None, "순이익 칸을 읽지 못했습니다"
    assert float(actual) == pytest.approx(expected, abs=0.5)


def test_every_summary_line_matches_python(recalculated, sample_ledger):
    """순이익만이 아니라 요약 16개 항목이 전부 맞아야 한다."""
    expected = compute(sample_ledger).as_summary()
    mismatched = []
    for key, row in L.SUMMARY.items():
        actual = recalculated.get(L.SHEET_LEDGER, f"B{row}")
        wanted = expected[key]
        if wanted is None:
            continue
        if actual is None or abs(float(actual) - float(wanted)) > 0.5:
            mismatched.append(f"{L.SUMMARY_LABELS[key][0]}: {actual!r} != {wanted!r}")
    assert mismatched == [], "; ".join(mismatched)


def test_recalculated_stock_matches_python(recalculated, sample_ledger):
    expected = compute(sample_ledger).products
    for index, product in enumerate(expected):
        cell = f"F{L.PRODUCT_FIRST_ROW + index}"
        actual = recalculated.get(L.SHEET_PRODUCTS, cell)
        assert float(actual) == pytest.approx(product.stock), product.label


def test_recalculated_product_table_matches_python(recalculated, sample_ledger):
    expected = compute(sample_ledger).products
    for index, product in enumerate(expected):
        row = L.PRODUCT_TABLE_FIRST_ROW + index
        assert float(recalculated.get(L.SHEET_LEDGER, f"C{row}")) == product.sold
        assert float(recalculated.get(L.SHEET_LEDGER, f"D{row}")) == product.revenue
        assert float(recalculated.get(L.SHEET_LEDGER, f"G{row}")) == product.profit
        assert float(recalculated.get(L.SHEET_LEDGER, f"I{row}")) == product.rank


def test_shipping_sheet_extracts_exactly_the_waiting_orders(recalculated, sample_ledger):
    waiting = [o for o in sample_ledger.orders if o.status == "대기"]
    for index, order in enumerate(waiting):
        row = 2 + index
        assert recalculated.get(L.SHEET_SHIPPING, f"A{row}") == order.customer
        assert recalculated.get(L.SHEET_SHIPPING, f"G{row}") == order.order_no
        assert float(recalculated.get(L.SHEET_SHIPPING, f"E{row}")) == order.qty
    after = recalculated.get(L.SHEET_SHIPPING, f"A{2 + len(waiting)}")
    assert after in ("", None), "대기 건이 끝난 줄은 비어 있어야 합니다"


def test_dashboard_stock_warning_lists_only_low_stock(recalculated, sample_ledger):
    totals = compute(sample_ledger)
    anchors = L.dash_rank_rows(min(sample_ledger.settings.days, L.MAX_DAY_ROWS))
    expected = [p.code for p in totals.low_stock_products]
    assert expected, "샘플에 재고 부족 상품이 있어야 검사가 의미 있다"
    for index, code in enumerate(expected):
        assert recalculated.get(L.SHEET_DASHBOARD, f"A{anchors['stock_first'] + index}") == code
    after = recalculated.get(L.SHEET_DASHBOARD, f"A{anchors['stock_first'] + len(expected)}")
    assert after in ("", None)


def test_dashboard_profit_ranking_is_sorted(recalculated, sample_ledger):
    totals = compute(sample_ledger)
    anchors = L.dash_rank_rows(min(sample_ledger.settings.days, L.MAX_DAY_ROWS))
    ordered = sorted(totals.products, key=lambda p: p.rank)
    for index, product in enumerate(ordered[:L.TOP_PRODUCTS]):
        row = anchors["rank_first"] + index
        assert recalculated.get(L.SHEET_DASHBOARD, f"B{row}") == product.code


def test_daily_revenue_matches_python(recalculated, sample_ledger):
    totals = compute(sample_ledger)
    for index, day in enumerate(totals.days):
        row = L.DASH["daily_first"] + index
        actual = recalculated.get(L.SHEET_DASHBOARD, f"B{row}")
        assert float(actual) == pytest.approx(day.revenue), str(day.day)


# ------------------------------------------------------------------ 되읽기
def test_ledger_survives_a_round_trip(built, sample_ledger):
    """import-orders 는 파일을 읽어 다시 만든다. 그 과정에서 값이 바뀌면 안 된다."""
    again = read_ledger(built)
    assert again.settings.name == sample_ledger.settings.name
    assert again.settings.start == sample_ledger.settings.start
    assert again.settings.payment_fee_rate == sample_ledger.settings.payment_fee_rate
    assert again.settings.vat_included == sample_ledger.settings.vat_included
    assert len(again.products) == len(sample_ledger.products)
    assert len(again.orders) == len(sample_ledger.orders)
    assert compute(again).profit == compute(sample_ledger).profit


def test_rebuilding_keeps_charts_and_validation(tmp_path, built, sample_ledger):
    """차트가 사라지면 두 번째 회차부터 대시보드가 빈 껍데기가 된다."""
    again = read_ledger(built)
    path = build_workbook(again, tmp_path / "rebuilt.xlsx")
    with zipfile.ZipFile(path) as archive:
        assert [n for n in archive.namelist() if n.startswith("xl/charts/chart")]
    sheet = load_workbook(path)[L.SHEET_ORDERS]
    assert sheet.data_validations.dataValidation


def test_reader_rejects_a_workbook_it_did_not_make(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "남의파일.xlsx"
    Workbook().save(path)
    with pytest.raises(LedgerFileError, match="시트"):
        read_ledger(path)


# ------------------------------------------------------------------ 리포트
@pytest.fixture(scope="session")
def report_data(built, sample_ledger):
    return gather(sample_ledger, built, use_recalc=True)


def test_report_numbers_come_from_the_spreadsheet(report_data):
    assert report_data.engine, "엑셀 재계산 결과를 써야 한다"
    assert report_data.mismatches == [], report_data.mismatches
    assert report_data.verified


def test_unknown_numbers_catches_invented_figures(report_data):
    allowed = {100.0, 20.0}
    assert unknown_numbers("이익은 100원입니다", allowed) == []
    assert unknown_numbers("매출이 987654원 늘었습니다", allowed) == [987654.0]


def test_insight_is_three_lines_using_only_real_numbers(report_data):
    lines, warnings = insight_lines(report_data, fake_ask, "dry-run")
    assert len(lines) == 3
    assert warnings == []
    assert banned_phrases.check(" ".join(lines)) == []


def test_insight_retries_when_the_model_invents_a_number(report_data):
    state = {"calls": 0}

    def liar(system, user, model="test", **_):
        state["calls"] += 1
        if state["calls"] == 1:
            return ("- 매출이 8,765,432원 늘었습니다\n"
                    "- 환불률이 좋습니다\n- 다음에도 잘 될 것입니다")
        return fake_ask(system, user, model)

    lines, warnings = insight_lines(report_data, liar, "test")
    assert state["calls"] == 2, "지어낸 숫자를 보면 다시 시켜야 한다"
    assert len(lines) == 3
    assert "8,765,432" not in " ".join(lines)


def test_insight_drops_lines_it_cannot_fix(report_data):
    def stubborn(system, user, model="test", **_):
        return "- 매출이 8,765,432원 늘었습니다\n- 둘째 줄\n- 셋째 줄"

    lines, warnings = insight_lines(report_data, stubborn, "test")
    assert all("8,765,432" not in line for line in lines)
    assert any("뺐습니다" in warning for warning in warnings)


def test_report_file_has_every_section(tmp_path, sample_ledger, report_data):
    lines, warnings = insight_lines(report_data, fake_ask, "dry-run")
    path = write_report(sample_ledger, report_data, tmp_path / "report.md",
                        lines, warnings, datetime(2026, 9, 16).astimezone())
    text = path.read_text(encoding="utf-8")
    for heading in ("## 정산 요약", "## 주문 현황", "## 상품별 실적",
                    "## 코멘트", "## 다음에 할 일"):
        assert heading in text, heading
    assert "재고 경고" in text
    assert "생성형 AI" in text, "AI 생성물 표시가 기본으로 들어가야 한다"
    assert banned_phrases.check(text) == []


def test_report_quotes_the_spreadsheet_profit(tmp_path, sample_ledger, report_data):
    lines, warnings = insight_lines(report_data, fake_ask, "dry-run")
    path = write_report(sample_ledger, report_data, tmp_path / "r.md",
                        lines, warnings, datetime(2026, 9, 16).astimezone())
    profit = compute(sample_ledger).profit
    assert f"{round(profit):,}원" in path.read_text(encoding="utf-8")


def test_report_warns_loudly_when_the_two_calculations_disagree(
        tmp_path, sample_ledger, report_data):
    broken = gather(sample_ledger, tmp_path / "없는파일.xlsx", use_recalc=False)
    broken.mismatches = ["순이익: 엑셀 1 vs 파이썬 2"]
    path = write_report(sample_ledger, broken, tmp_path / "warn.md",
                        ["한 줄"], [], datetime(2026, 9, 16).astimezone())
    text = path.read_text(encoding="utf-8")
    assert "⚠️" in text and "다릅니다" in text


def test_ai_label_can_be_turned_off(tmp_path, sample_ledger, report_data):
    path = write_report(sample_ledger, report_data, tmp_path / "plain.md",
                        ["한 줄"], [], datetime(2026, 9, 16).astimezone(), ai_label=False)
    assert "생성형 AI" not in path.read_text(encoding="utf-8")


def test_fake_ask_only_uses_numbers_from_the_facts_block(report_data):
    from groupbuy_ledger.report import _allowed_numbers, _facts

    text = fake_ask("", f"[수치]\n{_facts(report_data)}\n\n세 줄을 써라.")
    assert unknown_numbers(text, _allowed_numbers(report_data)) == []


def test_shipping_csv_has_only_waiting_orders(tmp_path, sample_ledger):
    path, count = write_shipping_csv(sample_ledger, tmp_path / "발송.csv")
    waiting = [o for o in sample_ledger.orders if o.status == "대기"]
    assert count == len(waiting)

    with path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert tuple(rows[0]) == SHIPPING_HEADER
    assert len(rows) == len(waiting) + 1
    assert [row[6] for row in rows[1:]] == [o.order_no for o in waiting]


def test_shipping_csv_keeps_korean_readable_in_excel(tmp_path, sample_ledger):
    """BOM 이 없으면 엑셀에서 한글이 깨져 보인다."""
    path, _ = write_shipping_csv(sample_ledger, tmp_path / "발송.csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")


# ------------------------------------------------------------------ CLI
def test_cli_runs_new_then_import_then_report(tmp_path):
    """완료 기준: new → import → report 순서 실행 성공."""
    out = tmp_path / "out"
    assert gb_cli.main([
        "new", "--name", "테스트 공구", "--products", str(SAMPLES / "products.csv"),
        "--start", "2026-09-01", "--end", "2026-09-14",
        "--cost", "9000", "--price", "15000", "--out", str(out),
    ]) == 0

    ledger_path = out / "groupbuy_테스트-공구.xlsx"
    assert ledger_path.is_file()

    assert gb_cli.main(["import-orders", str(ledger_path),
                        str(SAMPLES / "orders.csv")]) == 0
    assert len(read_ledger(ledger_path).orders) >= 20

    assert gb_cli.main(["report", str(ledger_path), "--dry-run",
                        "--out", str(out)]) == 0
    report = out / "report_테스트-공구.md"
    assert report.is_file()
    assert "정산 요약" in report.read_text(encoding="utf-8")
    assert (out / "발송목록_테스트-공구.csv").is_file()


def test_cli_demo_runs_the_whole_pipeline(tmp_path):
    """대시보드의 테스트 버튼이 누르는 명령. 세 단계를 한 번에 돈다."""
    out = tmp_path / "demo"
    assert gb_cli.main(["demo", "--name", "데모 공구", "--dry-run",
                        "--out", str(out)]) == 0
    assert (out / "groupbuy_데모-공구.xlsx").is_file()
    assert (out / "report_데모-공구.md").is_file()
    assert (out / "발송목록_데모-공구.csv").is_file()
    assert len(read_ledger(out / "groupbuy_데모-공구.xlsx").orders) >= 20


def test_cli_verify_passes_on_a_freshly_built_ledger(tmp_path):
    out = tmp_path / "v"
    gb_cli.main(["new", "--name", "검증 공구", "--products", str(SAMPLES / "products.csv"),
                 "--cost", "9000", "--price", "15000", "--out", str(out)])
    path = out / "groupbuy_검증-공구.xlsx"
    gb_cli.main(["import-orders", str(path), str(SAMPLES / "orders.csv")])
    assert gb_cli.main(["verify", str(path)]) == 0


def test_cli_import_returns_two_when_something_needs_a_human(tmp_path):
    out = tmp_path / "w"
    gb_cli.main(["new", "--name", "경고 공구", "--products", str(SAMPLES / "products.csv"),
                 "--cost", "9000", "--price", "15000", "--out", str(out)])
    path = out / "groupbuy_경고-공구.xlsx"
    stray = tmp_path / "stray.csv"
    stray.write_text("결제일,주문번호,상품코드,옵션,수량,결제금액,주문상태\n"
                     "2026-09-02,Q1,없는상품,특대,1,15000,완료\n", encoding="utf-8")
    assert gb_cli.main(["import-orders", str(path), str(stray)]) == 2


def test_cli_without_a_command_shows_help(capsys):
    assert gb_cli.main([]) == 1
    assert "new" in capsys.readouterr().out


# ------------------------------------------------------------------ 문서
def test_readme_covers_usage_and_is_clean():
    text = (BASE_DIR / "README.md").read_text(encoding="utf-8")
    for token in ("cli.py new", "import-orders", "report", "products.csv",
                  "orders.csv", "mapping.yaml", "INDEX", "SUMIFS"):
        assert token in text, token
    assert "XLOOKUP" in text, "쓰지 않는 함수와 그 이유를 적어야 한다"
    assert banned_phrases.check(text) == []


def test_product_manuals_are_detailed_and_clean():
    for name in ("admin.md", "client.md"):
        path = BASE_DIR / "docs" / name
        assert path.is_file(), name
        text = path.read_text(encoding="utf-8")
        assert len(text) > 4000, f"{name} 이 너무 짧습니다"
        assert banned_phrases.check(text) == []


def test_mapping_file_documents_how_to_extend_it():
    text = MAPPING.read_text(encoding="utf-8")
    assert "추가" in text
    mapping = load_mapping(MAPPING)
    assert set(mapping["columns"]) >= {"주문일", "주문번호", "상품코드", "수량", "결제금액"}
    assert set(mapping["status"]) == {"대기", "발송", "완료", "취소", "환불"}
