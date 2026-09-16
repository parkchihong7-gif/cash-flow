"""shared/ai_label.py, shared/banned_phrases.py 단위 테스트.

네트워크를 타지 않는다. shared/llm.py 는 실제 API 호출이 필요하므로
여기서는 순수 함수(strip_json_fence)만 검사한다.
"""

from __future__ import annotations

import pytest

from shared import ai_label, banned_phrases
from shared.llm import strip_json_fence


# ----------------------------------------------------------------- ai_label
def test_add_text_label_appends_korean_notice():
    result = ai_label.add_text_label("본문입니다.")
    assert result.endswith(ai_label.LABELS["ko"])
    assert result.startswith("본문입니다.")


def test_add_text_label_is_idempotent():
    once = ai_label.add_text_label("본문입니다.")
    twice = ai_label.add_text_label(once)
    assert once == twice
    assert once.count(ai_label.LABELS["ko"]) == 1


def test_add_text_label_english():
    result = ai_label.add_text_label("Body text.", lang="en")
    assert result.endswith(ai_label.LABELS["en"])


def test_add_text_label_unknown_lang_falls_back_to_korean():
    assert ai_label.add_text_label("본문", lang="fr").endswith(ai_label.LABELS["ko"])


def test_add_text_label_on_empty_text():
    assert ai_label.add_text_label("") == ai_label.LABELS["ko"]
    assert ai_label.add_text_label("   \n ") == ai_label.LABELS["ko"]


def _make_docx(path):
    import docx

    document = docx.Document()
    document.add_paragraph("테스트 문서")
    document.save(str(path))
    return path


def _make_pptx(path):
    import pptx

    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))
    return path


def _make_xlsx(path):
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "테스트"
    workbook.save(str(path))
    return path


def _read_marker(path):
    """저장된 파일에서 표시 문자열을 다시 읽는다."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        return docx.Document(str(path)).core_properties.comments
    if suffix == ".pptx":
        import pptx

        return pptx.Presentation(str(path)).core_properties.comments
    import openpyxl

    return openpyxl.load_workbook(str(path)).properties.description


@pytest.mark.parametrize(
    ("name", "builder"),
    [("doc.docx", _make_docx), ("deck.pptx", _make_pptx), ("sheet.xlsx", _make_xlsx)],
)
def test_add_metadata_writes_marker(tmp_path, name, builder):
    path = builder(tmp_path / name)

    marker = ai_label.add_metadata(path, tool="agency-kit", date="2026-09-16T00:00:00+00:00")

    assert marker == "AI-generated: true; tool: agency-kit; date: 2026-09-16T00:00:00+00:00"
    # python-docx/pptx 가 넣어둔 기본 comments 는 지우지 않고 뒤에 덧붙인다.
    assert marker in _read_marker(path)


def test_add_metadata_is_idempotent(tmp_path):
    path = _make_docx(tmp_path / "doc.docx")

    first = ai_label.add_metadata(path, tool="agency-kit", date="2026-09-16T00:00:00+00:00")
    second = ai_label.add_metadata(path, tool="agency-kit", date="2026-09-17T00:00:00+00:00")

    assert first == second
    assert _read_marker(path).count(ai_label.METADATA_PREFIX) == 1


def test_add_metadata_defaults_date_to_now(tmp_path):
    path = _make_docx(tmp_path / "doc.docx")

    marker = ai_label.add_metadata(path, tool="funnel-builder")

    assert marker.startswith("AI-generated: true; tool: funnel-builder; date: ")


def test_add_metadata_infers_tool_from_products_path(tmp_path, monkeypatch):
    products = tmp_path / "products"
    (products / "niche-research").mkdir(parents=True)
    monkeypatch.setattr(ai_label, "PRODUCTS_DIR", products)
    path = _make_docx(products / "niche-research" / "doc.docx")

    marker = ai_label.add_metadata(path)

    assert "tool: niche-research;" in marker


def test_add_metadata_rejects_unsupported_suffix(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("본문")

    with pytest.raises(ValueError, match="지원하지 않는 형식"):
        ai_label.add_metadata(path)


def test_add_metadata_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        ai_label.add_metadata(tmp_path / "없는파일.docx")


# ----------------------------------------------------------- banned_phrases
def test_banned_list_has_at_least_20_entries():
    assert len(banned_phrases.BANNED) >= 20
    assert len(set(banned_phrases.BANNED)) == len(banned_phrases.BANNED)


def test_check_finds_banned_phrase():
    assert "수익 보장" in banned_phrases.check("이 강의는 수익 보장 상품입니다.")


def test_check_ignores_spacing_variants():
    assert banned_phrases.check("수익보장!") == ["수익 보장"]
    assert banned_phrases.check("수익  보장") == ["수익 보장"]


def test_check_returns_multiple_in_banned_order():
    text = "무조건 성공하고 원금 보장까지 됩니다."
    found = banned_phrases.check(text)
    assert "원금 보장" in found
    assert "무조건" in found
    assert found == [p for p in banned_phrases.BANNED if p in found]


def test_check_clean_text():
    text = "반복 작업을 자동화해 시간을 절약합니다. 실제 성과는 검증이 필요합니다."
    assert banned_phrases.check(text) == []
    assert banned_phrases.is_clean(text)


def test_check_empty_text():
    assert banned_phrases.check("") == []


def test_assert_clean_passes_through():
    text = "반복 작업 자동화 도구입니다."
    assert banned_phrases.assert_clean(text) == text


def test_assert_clean_raises_with_found_list():
    with pytest.raises(banned_phrases.BannedPhraseError) as excinfo:
        banned_phrases.assert_clean("자동으로 돈이 들어옵니다.")
    assert "자동으로 돈" in excinfo.value.found


def test_project_readmes_are_clean():
    """저장소의 판매용 텍스트가 실제로 규칙을 지키는지 확인한다."""
    from pathlib import Path

    from shared.config import ROOT_DIR

    for readme in [ROOT_DIR / "README.md", *Path(ROOT_DIR / "products").rglob("README.md")]:
        assert banned_phrases.check(readme.read_text(encoding="utf-8")) == [], readme


# ---------------------------------------------------------------------- llm
def test_strip_json_fence_removes_json_fence():
    assert strip_json_fence('```json\n{"a": 1}\n```') == '{"a": 1}'


def test_strip_json_fence_removes_bare_fence():
    assert strip_json_fence('```\n{"a": 1}\n```') == '{"a": 1}'


def test_strip_json_fence_leaves_plain_json():
    assert strip_json_fence('  {"a": 1}  ') == '{"a": 1}'
