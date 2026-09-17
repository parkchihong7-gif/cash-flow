"""11번 자동화 대행 납품 키트 테스트 — 모듈 A·B·C + 납품 패키지.

모듈마다 지켜야 하는 것이 하나씩 있다.

    A 카카오 챗봇   **5초 안에 무언가를 돌려준다.** 못 돌려주면 고객 화면에
                   "오류가 발생했습니다" 가 뜬다
    B 주간 보고서   **Claude 가 숫자를 만들지 않는다.** 집계에 없는 숫자가 나오면
                   잡아낸다
    C 인스타 발행   **승인 없이는 아무것도 올라가지 않는다**

모듈 폴더 이름에 하이픈이 있어 import 할 수 없다. 파일 경로로 직접 싣는다.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from datetime import datetime, timedelta
from pathlib import Path
from types import ModuleType

import pytest

from shared import banned_phrases

BASE_DIR = Path(__file__).resolve().parents[1] / "products" / "agency-kit"
KAKAO = BASE_DIR / "kakao-faq-bot"
REPORT = BASE_DIR / "sheet-report"
INSTA = BASE_DIR / "insta-scheduler"
DELIVERABLES = BASE_DIR / "deliverables"


def _load(directory: Path, name: str) -> ModuleType:
    """모듈 폴더의 파일 하나를 고유한 이름으로 싣는다.

    `app`·`manage`·`run` 같은 이름은 상품끼리 겹친다. 경로를 지정해 싣고
    `agency_<폴더>_<파일>` 로 등록해 어느 테스트가 먼저 돌든 같게 만든다.
    """
    unique = f"agency_{directory.name.replace('-', '_')}_{name}"
    if unique in sys.modules:
        return sys.modules[unique]

    if str(directory) not in sys.path:
        sys.path.insert(0, str(directory))

    path = directory / f"{name}.py"
    spec = importlib.util.spec_from_file_location(unique, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[unique] = module
    spec.loader.exec_module(module)
    return module


# 모듈 A
kakao_app = _load(KAKAO, "app")
kakao_logs = _load(KAKAO, "logs")
kakao_matcher = _load(KAKAO, "matcher")
kakao_faq = _load(KAKAO, "faq_source")
kakao_manage = _load(KAKAO, "manage")
kakao_sample = _load(KAKAO, "sample_content")

# 모듈 B
report_aggregate = _load(REPORT, "aggregate")
report_narrate = _load(REPORT, "narrate")
report_writers = _load(REPORT, "writers")
report_notify = _load(REPORT, "notify")
report_run = _load(REPORT, "run")
report_workflow = _load(REPORT, "make_workflow")

# 모듈 C
insta_queue = _load(INSTA, "queue_store")
insta_publisher = _load(INSTA, "publisher")
insta_scheduler = _load(INSTA, "scheduler")
insta_captions = _load(INSTA, "captions")
insta_manage = _load(INSTA, "manage")

kit_cli = _load(BASE_DIR, "cli")


# ====================================================================== 모듈 A
def _payload(utterance: str, user_id: str = "kakao-test-1") -> dict:
    """오픈빌더가 실제로 보내는 모양."""
    return {
        "intent": {"id": "i1", "name": "블록"},
        "userRequest": {
            "timezone": "Asia/Seoul", "params": {},
            "block": {"id": "b1", "name": "폴백 블록"},
            "utterance": utterance, "lang": "ko",
            "user": {"id": user_id, "type": "accountId", "properties": {}},
        },
        "bot": {"id": "bot1", "name": "테스트봇"},
        "action": {"name": "faq", "clientExtra": {}, "params": {}, "id": "a1",
                   "detailParams": {}},
    }


@pytest.fixture
def kakao_client(tmp_path):
    from fastapi.testclient import TestClient

    book = kakao_faq.FaqBook(
        loader=lambda: kakao_faq.load_from_csv(KAKAO / "sample_faq.csv"))
    store = kakao_logs.LogStore(tmp_path / "logs.db")
    app = kakao_app.create_app(book=book, ask_fn=kakao_sample.fake_ask,
                               store=store, model="test")
    client = TestClient(app)
    client.store = store                       # 테스트에서 기록을 들여다본다
    return client


def test_skill_response_follows_the_openbuilder_spec(kakao_client):
    """규격을 어기면 고객 화면에 '오류가 발생했습니다' 만 뜬다."""
    body = kakao_client.post("/skill", json=_payload("영업시간이 어떻게 되나요?")).json()
    assert body["version"] == "2.0"
    outputs = body["template"]["outputs"]
    assert len(outputs) == 1
    assert isinstance(outputs[0]["simpleText"]["text"], str)


def test_a_matched_question_returns_the_sheet_answer_verbatim(kakao_client):
    """시트 문장은 고객이 검수한 문장이다. 한 글자도 고치지 않는다."""
    faqs = kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")
    expected = next(faq.answer for faq in faqs if "영업시간" in faq.question)

    text = kakao_client.post("/skill", json=_payload("영업시간이 어떻게 되나요?")).json()
    answer = text["template"]["outputs"][0]["simpleText"]["text"]
    assert answer.startswith(expected)


def test_an_unknown_question_gets_the_check_with_staff_tail(kakao_client):
    answer = kakao_client.post(
        "/skill", json=_payload("강아지 데려가도 되나요?")
    ).json()["template"]["outputs"][0]["simpleText"]["text"]
    assert "담당자" in answer


def test_the_answer_carries_the_ai_label(kakao_client):
    """인공지능기본법 제31조. 기본으로 켜져 있어야 한다."""
    answer = kakao_client.post(
        "/skill", json=_payload("주차 가능한가요?")
    ).json()["template"]["outputs"][0]["simpleText"]["text"]
    assert kakao_app.AI_LABEL_TEXT in answer


def test_a_slow_model_falls_back_inside_the_budget(tmp_path):
    """4초를 넘기면 버리고 폴백한다. 늦게 온 답은 쓸 데가 없다."""
    import time

    from fastapi.testclient import TestClient

    book = kakao_faq.FaqBook(
        loader=lambda: kakao_faq.load_from_csv(KAKAO / "sample_faq.csv"))
    app = kakao_app.create_app(
        book=book, ask_fn=kakao_sample.fake_ask_slow(10),
        store=kakao_logs.LogStore(tmp_path / "slow.db"), model="test", budget=0.6)
    client = TestClient(app)

    started = time.monotonic()
    answer = client.post(
        "/skill", json=_payload("이 질문은 시트에 없는 아주 낯선 질문입니다")
    ).json()["template"]["outputs"][0]["simpleText"]["text"]
    elapsed = time.monotonic() - started

    assert elapsed < 2.0, f"{elapsed:.1f}초나 걸렸습니다"
    assert "잠시 후" in answer


def test_a_broken_model_still_answers_in_spec(tmp_path):
    """어떤 예외가 나도 규격에 맞는 응답을 돌려준다."""
    from fastapi.testclient import TestClient

    def explode(*args, **kwargs):
        raise RuntimeError("모델이 터졌습니다")

    book = kakao_faq.FaqBook(
        loader=lambda: kakao_faq.load_from_csv(KAKAO / "sample_faq.csv"))
    client = TestClient(kakao_app.create_app(
        book=book, ask_fn=explode,
        store=kakao_logs.LogStore(tmp_path / "boom.db"), model="test"))

    body = client.post("/skill", json=_payload("처음 보는 질문입니다 정말로")).json()
    assert body["version"] == "2.0"
    assert kakao_app.FALLBACK_TEXT.splitlines()[0] in \
        body["template"]["outputs"][0]["simpleText"]["text"]


def test_an_empty_utterance_does_not_crash(kakao_client):
    body = kakao_client.post("/skill", json=_payload("")).json()
    assert body["version"] == "2.0"


def test_health_reports_the_faq_count(kakao_client):
    body = kakao_client.get("/health").json()
    assert body["status"] == "ok"
    assert body["faq_count"] >= 10
    assert body["ai_label"] is True


def test_the_sheet_is_cached_for_five_minutes():
    """질문마다 시트를 읽으면 그것만으로 1~2초를 쓴다."""
    calls = {"n": 0}

    def loader():
        calls["n"] += 1
        return kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")

    book = kakao_faq.FaqBook(loader=loader, cache_seconds=300)
    book.get(now=1000.0)
    book.get(now=1100.0)
    book.get(now=1290.0)
    assert calls["n"] == 1, "캐시 안에서 다시 읽었습니다"

    book.get(now=1400.0)
    assert calls["n"] == 2, "5분이 지났는데 다시 읽지 않았습니다"


def test_a_broken_sheet_keeps_serving_the_old_answers():
    """빈 답보다 예전 답이 낫다."""
    state = {"fail": False}

    def loader():
        if state["fail"]:
            raise RuntimeError("시트가 잠깐 안 됩니다")
        return kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")

    book = kakao_faq.FaqBook(loader=loader, cache_seconds=1)
    first = book.get(now=1000.0)
    state["fail"] = True
    again = book.get(now=2000.0)

    assert len(again) == len(first)
    assert "시트가 잠깐" in book.last_error


def test_columns_are_found_by_name_not_position(tmp_path):
    """고객이 칸 순서를 바꿔 놓는 일이 흔하다."""
    path = tmp_path / "faq.csv"
    path.write_text("카테고리,답변,수정일,질문\n기본,10시입니다,2026-01-01,몇 시에 여나요\n",
                    encoding="utf-8")
    faqs = kakao_faq.load_from_csv(path)
    assert faqs[0].question == "몇 시에 여나요"
    assert faqs[0].answer == "10시입니다"


def test_a_sheet_without_the_required_columns_is_refused(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("이름,메모\n가,나\n", encoding="utf-8")
    with pytest.raises(ValueError, match="질문"):
        kakao_faq.load_from_csv(path)


@pytest.mark.parametrize("question,needle", [
    ("영업시간이 어떻게 되나요?", "영업시간"),
    ("영업시간 어떻게 되나요", "영업시간"),
    ("환불 규정이 어떻게 되나요?", "환불"),
])
def test_literal_matching_catches_near_identical_questions(question, needle):
    """글자로 걸리면 Claude 를 안 부른다. 공짜이고 즉시 끝난다."""
    faqs = kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")
    index = kakao_matcher.literal_match(question, faqs)
    assert index is not None and needle in faqs[index].question


def test_literal_matching_does_not_guess():
    """비슷해 보여도 다른 질문이면 고르지 않는다."""
    faqs = kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")
    assert kakao_matcher.literal_match("강아지 데려가도 되나요?", faqs) is None


def test_an_out_of_range_index_is_treated_as_no_match():
    """모델이 목록에 없는 번호를 말하는 일이 실제로 있다."""
    faqs = kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")
    assert kakao_matcher.match("질문", faqs, lambda *a, **k: {"index": 999}, "m") is None
    assert kakao_matcher.match("질문", faqs, lambda *a, **k: {"index": 0}, "m") is None
    assert kakao_matcher.match("질문", faqs, lambda *a, **k: {"index": None}, "m") is None


def test_match_converts_one_based_numbers():
    faqs = kakao_faq.load_from_csv(KAKAO / "sample_faq.csv")
    assert kakao_matcher.match("질문", faqs, lambda *a, **k: {"index": 1}, "m") == 0


# ---------------------------------------------------------------- 기록·개인정보
@pytest.mark.parametrize("raw,gone", [
    ("제 번호 010-1234-5678 로 연락 주세요", "010-1234-5678"),
    ("메일은 hong@example.com 입니다", "hong@example.com"),
    ("카드 1234-5678-9012-3456 로 결제했어요", "1234-5678-9012-3456"),
    ("주민번호 900101-1234567 입니다", "900101-1234567"),
])
def test_personal_data_is_scrubbed_before_storing(raw, gone):
    """고객센터 질문에는 개인정보가 자주 딸려 온다."""
    cleaned = kakao_logs.scrub(raw)
    assert gone not in cleaned
    assert "[" in cleaned


def test_the_user_id_is_hashed_not_stored(tmp_path):
    store = kakao_logs.LogStore(tmp_path / "logs.db")
    store.add(user_id="kakao-real-user-id", question="영업시간", answer="10시",
              matched=1, elapsed_ms=12, source="sheet")

    text = (tmp_path / "logs.db").read_bytes().decode("utf-8", "replace")
    assert "kakao-real-user-id" not in text


def test_the_salt_changes_the_hash():
    first = kakao_logs.hash_user("u1", salt="salt-a")
    second = kakao_logs.hash_user("u1", salt="salt-b")
    assert first != second and len(first) == 16


def test_stats_collects_unmatched_questions(tmp_path):
    """못 답한 질문이 다음 달 리테이너의 내용물이 된다."""
    store = kakao_logs.LogStore(tmp_path / "logs.db")
    for _ in range(3):
        store.add(user_id="u", question="강아지 되나요", answer="모름",
                  matched=None, elapsed_ms=10, source="claude")
    store.add(user_id="u", question="영업시간", answer="10시",
              matched=1, elapsed_ms=10, source="sheet")

    rows = store.unmatched()
    assert rows[0]["question"] == "강아지 되나요" and rows[0]["hits"] == 3
    assert store.totals()["missed"] == 3


def test_purge_removes_old_records(tmp_path):
    """계약서에 적은 보관 기간을 지키는 손잡이."""
    store = kakao_logs.LogStore(tmp_path / "logs.db")
    store.add(user_id="u", question="q", answer="a", matched=1, elapsed_ms=5)
    assert store.purge_before("2999-01-01") == 1
    assert store.totals()["total"] == 0


def test_manage_stats_runs_on_an_empty_database(tmp_path, capsys):
    assert kakao_manage.main(["--db", str(tmp_path / "x.db"), "stats"]) == 0
    assert "기록이 없습니다" in capsys.readouterr().out


# ====================================================================== 모듈 B
@pytest.fixture(scope="module")
def sales():
    return report_aggregate.load_csv(REPORT / "sample_sales.csv")


def test_the_sample_sheet_aggregates(sales):
    result = report_aggregate.summarize(sales, period="week")
    assert result.total > 0 and result.count > 0
    assert result.prev_total > 0
    assert len(result.top) == 5
    assert result.group_column == "채널"


def test_totals_match_a_hand_count(sales):
    """집계가 틀리면 보고서 전체를 못 믿는다. 따로 세어 맞춰 본다."""
    result = report_aggregate.summarize(sales, period="week")
    window = sales.copy()
    window["_d"] = window["날짜"].astype(str)
    hand = int(window[(window["_d"] >= result.start) & (window["_d"] <= result.end)]
               ["매출"].sum())
    assert result.total == hand


def test_columns_are_found_by_any_common_name(tmp_path):
    import pandas as pd

    frame = pd.DataFrame({
        "일자": ["2026-09-01", "2026-09-02"],
        "결제금액": ["12,000원", "8,000원"],
        "구분": ["온라인", "오프라인"],
    })
    result = report_aggregate.summarize(frame, period="week", end_date="2026-09-02")
    assert result.total == 20000, "'12,000원' 같은 글자도 숫자로 읽어야 합니다"
    assert result.group_column == "구분"


def test_a_missing_date_column_says_what_is_needed():
    import pandas as pd

    with pytest.raises(report_aggregate.ColumnMissing, match="날짜"):
        report_aggregate.summarize(pd.DataFrame({"금액": [1]}), period="week")


def test_rows_without_a_readable_date_are_dropped_and_reported():
    import pandas as pd

    frame = pd.DataFrame({
        "날짜": ["2026-09-01", "언젠가", "2026-09-02"],
        "매출": [1000, 9999, 2000],
    })
    result = report_aggregate.summarize(frame, period="week", end_date="2026-09-02")
    assert result.total == 3000
    assert any("읽지 못한" in note for note in result.notes)


def test_narration_rejects_invented_numbers(sales):
    """Claude 가 더한 숫자는 사람이 검산할 수 없다. 그래서 막는다."""
    aggregate = report_aggregate.summarize(sales, period="week")

    def liar(system, user, **kwargs):
        return {"reading": ["매출이 987654321원 늘었습니다."],
                "actions": ["뭔가 하기", "또 뭔가 하기"], "caution": ""}

    narration = report_narrate.narrate(aggregate, liar, "test")
    assert not narration.verified
    assert any("집계에 없는 숫자" in w or "대조하지 못했" in w for w in narration.warnings)


def test_narration_accepts_quoted_numbers(sales):
    aggregate = report_aggregate.summarize(sales, period="week")

    def honest(system, user, **kwargs):
        return {
            "reading": [f"이번 기간 합계는 {aggregate.total:,}원입니다.",
                        "직전 기간보다 늘었습니다.", "건수를 함께 보세요."],
            "actions": ["월요일에 재고 확인", "미결제 고객에게 안내"],
            "caution": "현금 결제는 빠져 있을 수 있습니다.",
        }

    narration = report_narrate.narrate(aggregate, honest, "test")
    assert narration.verified and narration.tries == 1


def test_narration_retries_once_before_giving_up(sales):
    aggregate = report_aggregate.summarize(sales, period="week")
    calls = {"n": 0}

    def flaky(system, user, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            return {"reading": ["매출이 123456789원입니다."], "actions": [], "caution": ""}
        return {"reading": [f"합계는 {aggregate.total:,}원입니다."],
                "actions": ["할 일 하나", "할 일 둘"], "caution": ""}

    narration = report_narrate.narrate(aggregate, flaky, "test")
    assert calls["n"] == 2 and narration.verified


def test_small_numbers_are_not_flagged():
    """'3줄', '2가지' 같은 수는 순위·개수라 대조하지 않는다."""
    assert report_narrate.unknown_numbers("할 일 2가지를 3일 안에", {"100"}) == []


def test_report_files_are_written(tmp_path, sales):
    aggregate = report_aggregate.summarize(sales, period="week")
    narration = report_run._fake_narration(aggregate)
    md_path, docx_path = report_writers.write_all(
        aggregate, narration, tmp_path, title="테스트 보고서", stamp="20260101")

    assert md_path.is_file() and docx_path.is_file()
    assert docx_path.stat().st_size > 5000

    text = md_path.read_text(encoding="utf-8")
    assert f"{aggregate.total:,}원" in text
    assert "생성형 AI" in text, "AI 생성물 표시가 빠졌습니다"


def test_the_docx_opens_and_has_tables(tmp_path, sales):
    from docx import Document

    aggregate = report_aggregate.summarize(sales, period="week")
    narration = report_run._fake_narration(aggregate)
    _, docx_path = report_writers.write_all(aggregate, narration, tmp_path,
                                            stamp="20260102")
    document = Document(docx_path)
    assert len(document.tables) >= 1
    assert any("보고서" in paragraph.text for paragraph in document.paragraphs)


def test_run_end_to_end_with_the_sample_csv(tmp_path):
    code = report_run.main(["--csv", str(REPORT / "sample_sales.csv"),
                            "--period", "week", "--dry-run", "--out", str(tmp_path)])
    assert code in (0, 2)
    assert list(tmp_path.glob("report_*.md")) and list(tmp_path.glob("report_*.docx"))


def test_slack_without_a_webhook_is_not_an_error():
    result = report_notify.send_slack("보고서", webhook="")
    assert not result.sent and "선택 기능" in result.detail


def test_slack_posts_the_text_when_configured():
    seen = {}

    class _Response:
        status = 200

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    def opener(request, timeout=None):
        seen["url"] = request.full_url
        seen["body"] = json.loads(request.data.decode("utf-8"))
        return _Response()

    result = report_notify.send_slack("주간 보고서입니다",
                                      webhook="https://hooks.slack.test/x",
                                      opener=opener)
    assert result.sent and seen["body"]["text"] == "주간 보고서입니다"


def test_email_without_settings_is_not_an_error(monkeypatch):
    monkeypatch.delenv("SMTP_HOST", raising=False)
    monkeypatch.delenv("REPORT_EMAIL_TO", raising=False)
    result = report_notify.send_email("제목", "본문")
    assert not result.sent and "선택 기능" in result.detail


# ------------------------------------------------------------- n8n 워크플로
def test_workflow_reuses_the_n8n_node_templates():
    """노드 규격을 두 군데 적어 두면 한쪽이 조용히 낡는다."""
    workflow = report_workflow.build()
    schedule = report_workflow.load_template("schedule-trigger")
    assert workflow["nodes"][0]["type"] == schedule["n8n_type"]
    assert workflow["nodes"][0]["typeVersion"] == schedule["type_version"]


def test_workflow_runs_weekly_and_calls_the_server():
    workflow = report_workflow.build(hour=9, minute=0, weekday=1)
    cron = workflow["nodes"][0]["parameters"]["rule"]["interval"][0]["expression"]
    assert cron == "0 9 * * 1"
    assert workflow["nodes"][1]["parameters"]["method"] == "POST"
    assert len(workflow["connections"]) == 2


def test_workflow_holds_no_real_credentials():
    """n8n credentials 는 이름만 둔다 (CLAUDE.md §7)."""
    text = json.dumps(report_workflow.build(), ensure_ascii=False)
    assert "xoxb-" not in text and "sk-ant" not in text
    assert "자격증명" in text


def test_the_shipped_workflow_file_matches_the_builder():
    """workflow.json 이 손으로 고쳐져 낡지 않았는지 본다."""
    shipped = json.loads((REPORT / "workflow.json").read_text(encoding="utf-8"))
    assert [node["type"] for node in shipped["nodes"]] == \
           [node["type"] for node in report_workflow.build()["nodes"]]


def test_serve_returns_a_report():
    from fastapi.testclient import TestClient

    serve = _load(REPORT, "serve")
    client = TestClient(serve.app)
    body = client.post("/run", json={"period": "week", "dry_run": True}).json()
    assert body["ok"] and body["total"] > 0
    assert "매출" in body["summary"]


# ====================================================================== 모듈 C
@pytest.fixture
def queue(tmp_path):
    return insta_queue.QueueStore(tmp_path / "queue.db")


def _fake_ig(published: list):
    class _Client:
        def publish_post(self, media_url, caption, media_type="IMAGE"):
            published.append((media_url, caption, media_type))
            return insta_publisher.PublishResult(ok=True, media_id="ig-123")

    return _Client()


def test_nothing_publishes_without_approval(queue):
    """이 테스트가 이 모듈에서 가장 중요하다."""
    published: list = []
    queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="예시입니다")

    runner = insta_scheduler.Runner(store=queue, client=_fake_ig(published))
    result = runner.tick()

    assert published == [], "승인하지 않은 글이 발행되었습니다"
    assert result.published == []


def test_an_approved_post_publishes_when_due(queue):
    published: list = []
    post_id = queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="예시입니다",
                        hashtags="#태그")
    queue.approve(post_id, "홍길동")

    result = insta_scheduler.Runner(store=queue, client=_fake_ig(published)).tick()

    assert result.published == [post_id]
    assert published[0][1] == "예시입니다\n\n#태그"
    assert queue.get(post_id).status == "published"


def test_a_future_post_waits(queue):
    published: list = []
    future = (datetime.now() + timedelta(days=2)).strftime("%Y-%m-%d %H:%M")
    post_id = queue.add(future, "https://x/1.jpg", caption="나중에")
    queue.approve(post_id, "홍길동")

    assert insta_scheduler.Runner(store=queue, client=_fake_ig(published)).tick().published == []


def test_editing_the_caption_revokes_approval(queue):
    """승인한 문장과 다른 글이 나가면 안 된다."""
    post_id = queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="처음 문장")
    queue.approve(post_id, "홍길동")
    assert queue.get(post_id).status == "approved"

    queue.set_caption(post_id, "고친 문장")
    assert queue.get(post_id).status == "draft"
    assert queue.get(post_id).approved_by == ""


def test_approval_needs_a_name(queue):
    post_id = queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="예시")
    with pytest.raises(ValueError, match="승인자"):
        queue.approve(post_id, "  ")


def test_an_empty_caption_cannot_be_approved(queue):
    post_id = queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="")
    with pytest.raises(ValueError, match="캡션"):
        queue.approve(post_id, "홍길동")


def test_a_failure_keeps_the_approval(queue):
    """고치고 다시 돌리면 그대로 나가야 한다."""
    class _Broken:
        def publish_post(self, *args, **kwargs):
            raise insta_publisher.InstagramError("[HTTP 190] 토큰 만료")

    post_id = queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="예시")
    queue.approve(post_id, "홍길동")

    result = insta_scheduler.Runner(store=queue, client=_Broken()).tick()
    assert result.failed and queue.get(post_id).status == "failed"
    assert "190" in queue.get(post_id).last_error

    queue.retry(post_id)
    assert queue.get(post_id).status == "approved"


def test_the_daily_limit_is_respected(queue):
    """인스타는 24시간에 25건까지만 받는다."""
    published: list = []
    for index in range(3):
        post_id = queue.add("2020-01-01 09:00", f"https://x/{index}.jpg", caption="예시")
        queue.approve(post_id, "홍길동")

    runner = insta_scheduler.Runner(store=queue, client=_fake_ig(published), daily_limit=2)
    result = runner.tick()

    assert len(result.published) == 2
    assert any("한도" in reason for _, reason in result.skipped)


def test_csv_import_creates_drafts(queue):
    added = queue.import_csv(INSTA / "queue.csv")
    assert len(added) == 3
    assert all(post.status == "draft" for post in queue.all())
    # 같은 줄을 다시 넣지 않는다
    assert queue.import_csv(INSTA / "queue.csv") == []


def test_pending_shows_posts_waiting_for_a_human(queue):
    queue.add("2020-01-01 09:00", "https://x/1.jpg", caption="승인 안 된 글")
    runner = insta_scheduler.Runner(store=queue, client=None, dry_run=True)
    assert [post.id for post in runner.pending()] == [1]


# ------------------------------------------------------------------ 발행 규격
class _FakeResponse:
    def __init__(self, payload: dict) -> None:
        self._body = json.dumps(payload).encode("utf-8")

    def read(self) -> bytes:
        return self._body

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def test_publishing_uses_the_two_step_container_flow():
    """컨테이너를 만들고 → 발행한다. 순서가 규격이다."""
    seen: list = []

    def opener(request, timeout=None):
        seen.append(request.full_url)
        if "media_publish" in request.full_url:
            return _FakeResponse({"id": "media-1"})
        return _FakeResponse({"id": "container-1"})

    client = insta_publisher.InstagramClient(access_token="t", ig_user_id="17841")
    client.opener = opener
    result = client.publish_post("https://x/1.jpg", "캡션")

    assert len(seen) == 2
    assert seen[0].endswith("17841/media")
    assert seen[1].endswith("17841/media_publish")
    assert result.media_id == "media-1"


def test_reels_wait_for_processing():
    """릴스는 처리 시간이 걸린다. 다 될 때까지 기다린다."""
    states = ["IN_PROGRESS", "IN_PROGRESS", "FINISHED"]
    slept: list = []

    def opener(request, timeout=None):
        url = request.full_url
        if "media_publish" in url:
            return _FakeResponse({"id": "media-2"})
        if "status_code" in url:
            return _FakeResponse({"status_code": states.pop(0)})
        return _FakeResponse({"id": "container-2"})

    client = insta_publisher.InstagramClient(access_token="t", ig_user_id="17841")
    client.opener = opener
    client.sleeper = slept.append
    result = client.publish_post("https://x/v.mp4", "캡션", media_type="REELS")

    assert result.media_id == "media-2"
    assert len(slept) == 2, "처리 중인데 기다리지 않았습니다"


def test_a_processing_error_is_explained():
    def opener(request, timeout=None):
        if "status_code" in request.full_url:
            return _FakeResponse({"status_code": "ERROR"})
        return _FakeResponse({"id": "container-3"})

    client = insta_publisher.InstagramClient(access_token="t", ig_user_id="17841")
    client.opener = opener
    with pytest.raises(insta_publisher.InstagramError, match="영상 규격"):
        client.publish_post("https://x/v.mp4", "캡션", media_type="REELS")


def test_token_errors_tell_what_to_do():
    import urllib.error

    def opener(request, timeout=None):
        raise urllib.error.HTTPError(
            request.full_url, 400, "Bad Request", {},
            __import__("io").BytesIO(json.dumps(
                {"error": {"code": 190, "message": "expired"}}).encode()))

    client = insta_publisher.InstagramClient(access_token="t", ig_user_id="1")
    client.opener = opener
    with pytest.raises(insta_publisher.InstagramError, match="refresh-token"):
        client.create_container("https://x/1.jpg", "캡션")


def test_there_are_no_follow_or_like_functions():
    """CLAUDE.md §3-4. 만들지 않은 것을 테스트로 고정한다."""
    source = (INSTA / "publisher.py").read_text(encoding="utf-8")
    forbidden = ("def follow", "def unfollow", "def like(", "def send_dm",
                 "/followers", "/likes", "def dm")
    for needle in forbidden:
        assert needle not in source, f"금지된 기능이 생겼습니다: {needle}"


# --------------------------------------------------------------------- 캡션
def test_caption_draft_cleans_hashtags():
    caption, hashtags, warnings = insta_captions.draft_caption(
        "가을 신상 사진", lambda *a, **k: {
            "caption": "오늘 준비한 것들입니다.",
            "hashtags": ["동네가게", "#동네가게", "#가을", "", "#신상"]},
        "test")
    assert caption
    assert hashtags == "#동네가게 #가을 #신상", "중복·빈 태그를 걸러야 합니다"
    assert warnings == []


def test_too_many_hashtags_are_trimmed():
    tags = insta_captions.clean_hashtags([f"#태그{n}" for n in range(30)])
    assert len(tags.split()) == insta_captions.MAX_TAGS


def test_begging_wording_is_flagged():
    """'맞팔해요' 는 정책상 위험하고 도달에도 나쁘다."""
    _, _, warnings = insta_captions.draft_caption(
        "메모", lambda *a, **k: {"caption": "팔로우 부탁드려요! 맞팔해요",
                                 "hashtags": []}, "test")
    assert any("부탁조" in w for w in warnings)


def test_banned_phrases_in_a_caption_are_flagged():
    _, _, warnings = insta_captions.draft_caption(
        "메모", lambda *a, **k: {"caption": "이거 하나면 수익 보장 됩니다",
                                 "hashtags": []}, "test")
    assert any("쓰면 안 되는" in w for w in warnings)


def test_a_sponsored_post_must_say_so():
    _, _, warnings = insta_captions.draft_caption(
        "메모", lambda *a, **k: {"caption": "좋은 제품 소개합니다", "hashtags": []},
        "test", sponsored=True)
    assert any("광고" in w for w in warnings)


# =============================================================== 키트 전체
def test_doctor_reports_the_environment(capsys):
    code = kit_cli.main(["doctor"])
    out = capsys.readouterr().out
    assert code in (0, 2)
    assert "fastapi" in out and "deliverables/INSTALL_GUIDE.md" in out


def test_demo_runs_all_three_modules(tmp_path, capsys):
    code = kit_cli.main(["demo", "--dry-run", "--out", str(tmp_path)])
    out = capsys.readouterr().out
    assert code == 0, out
    assert "모듈 A" in out and "모듈 B" in out and "모듈 C" in out

    folder = next(tmp_path.iterdir())
    for name in ("demo_summary.md", "kakao_demo.json", "insta_demo.md"):
        assert (folder / name).is_file(), name
    assert list(folder.glob("report_*.docx"))


def test_the_package_leaves_out_secrets(tmp_path, capsys):
    assert kit_cli.main(["package", "--client", "테스트가게", "--plan", "premium",
                         "--out", str(tmp_path)]) == 0
    folder = next(tmp_path.iterdir())

    for name in kit_cli.PACKAGE_FILES:
        assert (folder / name).is_file(), name
    assert (folder / "kakao-faq-bot" / "app.py").is_file()

    leaked = [str(path) for path in folder.rglob("*")
              if path.name in {".env", "credentials.json"} or path.suffix == ".db"]
    assert leaked == [], f"납품물에 비밀이 들어갔습니다: {leaked}"


def test_the_basic_package_only_ships_the_chatbot(tmp_path):
    kit_cli.main(["package", "--client", "가게", "--plan", "basic", "--out", str(tmp_path)])
    folder = next(tmp_path.iterdir())
    assert (folder / "kakao-faq-bot").is_dir()
    assert not (folder / "sheet-report").exists()


def test_the_package_fills_in_the_client_name(tmp_path):
    kit_cli.main(["package", "--client", "동네빵집", "--plan", "standard",
                  "--out", str(tmp_path)])
    folder = next(tmp_path.iterdir())
    checklist = (folder / "HANDOVER_CHECKLIST.md").read_text(encoding="utf-8")
    assert "고객 상호: 동네빵집" in checklist


# --------------------------------------------------------------------- 문서
@pytest.mark.parametrize("name", [
    "INSTALL_GUIDE.md", "RETAINER_CONTRACT_TEMPLATE.md",
    "HANDOVER_CHECKLIST.md", "KMONG_LISTING.md", "AI_NOTICE.md",
])
def test_deliverables_exist_and_are_clean(name):
    path = DELIVERABLES / name
    assert path.is_file(), name
    text = path.read_text(encoding="utf-8")
    assert len(text) > 1500, f"{name} 이 너무 짧습니다"
    assert banned_phrases.check(text) == [], name


def test_the_kmong_listing_has_three_packages():
    text = (DELIVERABLES / "KMONG_LISTING.md").read_text(encoding="utf-8")
    for tier in ("BASIC", "STANDARD", "PREMIUM"):
        assert tier in text
    assert "550,000원" in text and "880,000원" in text and "1,200,000원" in text


def test_the_checklist_really_has_twenty_items():
    text = (DELIVERABLES / "HANDOVER_CHECKLIST.md").read_text(encoding="utf-8")
    assert text.count("- [ ] **") == 20


def test_the_contract_leaves_money_blank():
    """금액은 협의해서 채우는 자리다. 미리 박아 두면 안 된다."""
    text = (DELIVERABLES / "RETAINER_CONTRACT_TEMPLATE.md").read_text(encoding="utf-8")
    assert "월 유지보수료: **[        ]원**" in text
    assert "초안입니다" in text


def test_the_ai_notice_cites_the_law():
    text = (DELIVERABLES / "AI_NOTICE.md").read_text(encoding="utf-8")
    assert "인공지능기본법" in text and "3,000만 원" in text


def test_module_readmes_explain_what_is_missing():
    """인스타 README 는 왜 팔로우 기능이 없는지 밝혀야 한다."""
    text = (INSTA / "README.md").read_text(encoding="utf-8")
    assert "Meta" in text and "정책" in text
    assert "자동 팔로우" in text


def test_the_kit_readme_says_it_does_not_promise_sales():
    text = (BASE_DIR / "README.md").read_text(encoding="utf-8")
    assert "매출을 보장하지 않" in text or "매출이나 문의 증가를 약속하지 않" in text
    assert banned_phrases.check(text) == []


def test_product_manuals_are_detailed_and_clean():
    for name in ("admin.md", "client.md"):
        path = BASE_DIR / "docs" / name
        assert path.is_file(), name
        text = path.read_text(encoding="utf-8")
        assert len(text) > 4000, f"{name} 이 너무 짧습니다"
        assert banned_phrases.check(text) == []


def test_env_examples_carry_no_real_keys():
    for path in BASE_DIR.rglob(".env.example"):
        text = path.read_text(encoding="utf-8")
        assert "sk-ant-" not in text, path
        for line in text.splitlines():
            if "=" in line and not line.startswith("#"):
                key, _, value = line.partition("=")
                if key.strip().endswith(("KEY", "TOKEN", "SECRET", "PASSWORD")):
                    assert value.strip() == "", f"{path}: {key} 에 값이 들어 있습니다"
