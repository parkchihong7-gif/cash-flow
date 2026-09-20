"""16번 해외 연사 초청 관리 테스트.

핵심 검사 셋:
1. **무비자 + 강연료** 조합에 경고가 뜨는가 (연사가 공항에서 돌아가는 경우)
2. **그로스업 계산**이 맞는가 (141만 원이 예산에서 빠지는 경우)
3. **판정하지 않는가** — 모든 산출물에 확인하라는 문구가 있는가
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from conftest import load_product_cli
from core.registry import Registry
from speaker_desk.brief import write_brief
from speaker_desk.budget import build as build_budget, write_xlsx
from speaker_desk.checklist import TEMPLATE, build as build_checklist, overdue_count
from speaker_desk.letters import NOT_LEGAL_ADVICE, contract_text, invitation_text
from speaker_desk.roster import Event, RosterError, Speaker, load_event
from speaker_desk.tax import (
    BASE_RATE, NOT_TAX_ADVICE, TOTAL_RATE, TREATY_DOCS, compute, gross_up,
)
from speaker_desk.visa import MUST_CONFIRM, SHORT_STAY_DAYS, assess

PRODUCT = Path(__file__).resolve().parents[1] / "products" / "speaker-desk"
EVENT_DAY = date(2026, 11, 20)


def _speaker(**kwargs) -> Speaker:
    base = dict(name="Jane Doe", country="United States", fee_krw=5_000_000,
                arrival=date(2026, 11, 18), departure=date(2026, 11, 22),
                visa_waiver=True, utc_offset=-5)
    base.update(kwargs)
    return Speaker(**base)


def _event(**kwargs) -> Event:
    base = dict(title="국제 AI 컨퍼런스 2026", event_date=EVENT_DAY,
                venue="코엑스", host="한국인공지능학회",
                speakers=[_speaker()])
    base.update(kwargs)
    return Event(**base)


# --------------------------------------- 무비자 + 강연료 = 가장 흔한 사고
def test_visa_waiver_plus_a_fee_raises_a_warning():
    """사증면제 대상국이라도 대가를 받으면 무비자로 강연할 수 없다."""
    hint = assess(paid=True, days=5, visa_waiver=True)
    assert hint.kind == "C-4"
    assert hint.warnings
    assert "무비자로 강연할 수 없습니다" in hint.warnings[0]


def test_unpaid_visit_from_a_waiver_country_has_no_warning():
    hint = assess(paid=False, days=3, visa_waiver=True)
    assert hint.kind == "무비자/K-ETA"
    assert not hint.warnings


def test_unpaid_visit_without_waiver_points_at_c3():
    assert assess(paid=False, days=3, visa_waiver=False).kind == "C-3"


def test_expenses_only_is_flagged_as_a_judgement_call():
    """실비 지원을 대가로 볼지는 사안마다 다르다. 단정하면 안 된다."""
    hint = assess(paid=False, days=4, visa_waiver=True, expenses_only=True)
    assert hint.warnings
    assert "사안마다 다릅니다" in hint.warnings[0]


def test_over_ninety_days_needs_another_status():
    hint = assess(paid=True, days=SHORT_STAY_DAYS + 1, visa_waiver=False)
    assert hint.kind == "장기"
    assert "E-1" in hint.warnings[0] or "E-7" in hint.warnings[0]


def test_the_tool_never_claims_to_decide():
    assert "판정이 아닙니다" in MUST_CONFIRM
    assert "출입국" in MUST_CONFIRM


# ------------------------------------------------------------ 원천징수
def test_default_rate_is_twenty_two_percent():
    """소득세 20% + 지방소득세 2%."""
    assert BASE_RATE == 0.20
    assert TOTAL_RATE == pytest.approx(0.22)


def test_gross_basis_deducts_from_the_contract_amount():
    result = compute(5_000_000, "gross")
    assert result.gross == 5_000_000
    assert result.tax == 1_100_000
    assert result.net == 3_900_000


def test_net_basis_grosses_up_and_costs_the_host_more():
    """500만 원을 손에 쥐어 드리려면 641만 원을 잡아야 한다."""
    result = compute(5_000_000, "net")
    assert result.net == 5_000_000
    assert result.gross == 6_410_256
    assert result.extra_for_net == 1_410_256


def test_gross_up_is_the_inverse_of_the_deduction():
    for amount in (1_000_000, 3_300_000, 12_345_678):
        grossed = gross_up(amount)
        assert compute(grossed, "gross").net == pytest.approx(amount, abs=1)


def test_treaty_rate_is_applied_and_marked():
    result = compute(5_000_000, "gross", treaty_rate=0.0)
    assert result.tax == 0
    assert result.net == 5_000_000
    assert result.treaty


def test_treaty_docs_name_the_certificate_of_residence():
    joined = " ".join(TREATY_DOCS)
    assert "거주자증명서" in joined
    assert "Certificate of Residence" in joined


def test_tax_module_says_it_is_not_advice():
    assert "세무 자문이 아닙니다" in NOT_TAX_ADVICE


def test_bad_inputs_are_refused():
    with pytest.raises(ValueError):
        compute(-1)
    with pytest.raises(ValueError, match="gross"):
        compute(100, "무언가")
    with pytest.raises(ValueError):
        compute(100, treaty_rate=1.5)


# ---------------------------------------------------------------- 명부
def test_a_fee_and_expenses_only_cannot_both_be_true():
    """비자 판단이 달라지는 값이라 둘 다일 수 없다."""
    with pytest.raises(Exception, match="둘 다일 수 없습니다"):
        _speaker(fee_krw=1_000_000, expenses_only=True)


def test_departure_before_arrival_is_refused():
    with pytest.raises(Exception, match="출국일이 입국일보다"):
        _speaker(arrival=date(2026, 11, 22), departure=date(2026, 11, 18))


def test_duplicate_speaker_names_are_refused():
    with pytest.raises(Exception, match="이름이 겹칩니다"):
        _event(speakers=[_speaker(), _speaker()])


def test_stay_days_counts_both_ends():
    assert _speaker().stay_days == 5


def test_jetlag_note_warns_about_the_morning_after():
    note = _speaker(utc_offset=-5).jetlag_note     # 시차 14시간
    assert "도착 다음 날 오전" in note
    assert "시차가 없습니다" in _speaker(utc_offset=9).jetlag_note


def test_sample_event_file_loads():
    event = load_event(PRODUCT / "event.yaml")
    assert len(event.speakers) == 2
    assert event.speaker("Jane Doe").paid


def test_unknown_speaker_lists_the_known_ones():
    with pytest.raises(RosterError, match="있는 연사"):
        _event().speaker("없는 사람")


# ------------------------------------------------------------ 체크리스트
def test_certificate_of_residence_is_requested_at_d80():
    """나라에 따라 발급에 몇 주 걸린다. 늦게 요청하면 조약을 못 쓴다."""
    spec = next(item for item in TEMPLATE if "거주자증명서" in item.title)
    assert spec.days_before == 80
    assert spec.hard
    assert spec.only_paid


def test_unpaid_speakers_skip_the_tax_tasks():
    event = _event(speakers=[_speaker(fee_krw=0, expenses_only=True)])
    titles = [task.title for task in build_checklist(event, event.speakers[0])]
    assert not any("거주자증명서" in title for title in titles)
    assert not any("원천징수" in title for title in titles)
    assert any("비자 신청" in title for title in titles)


def test_interpreter_tasks_appear_only_when_needed():
    event = _event(speakers=[_speaker(needs_interpreter=True)])
    with_interp = [t.title for t in build_checklist(event, event.speakers[0])]
    assert any("통역사 섭외" in title for title in with_interp)

    plain = _event(speakers=[_speaker(name="B", needs_interpreter=False)])
    without = [t.title for t in build_checklist(plain, plain.speakers[0])]
    assert not any("통역사 섭외" in title for title in without)


def test_overdue_counts_only_the_unrecoverable_ones():
    event = _event()
    tasks = build_checklist(event, event.speakers[0])
    # 행사 30일 전이면 D-120~D-45 가 다 지나 있다.
    late = overdue_count(tasks, EVENT_DAY - __import__("datetime").timedelta(days=30))
    assert late > 0
    soft_past = [t for t in tasks if not t.hard
                 and t.due < EVENT_DAY - __import__("datetime").timedelta(days=30)]
    assert soft_past, "필수가 아닌 지난 항목도 있어야 한다"


def test_not_everything_is_marked_required():
    """전부 필수면 아무것도 필수가 아니게 된다."""
    hard = [item for item in TEMPLATE if item.hard]
    assert 5 <= len(hard) <= 12, "필수 항목이 너무 많거나 적습니다"


def test_dday_labels_read_naturally():
    event = _event()
    tasks = {task.title: task for task in build_checklist(event, event.speakers[0])}
    assert tasks["행사 당일"].dday == "D-day"
    assert tasks["비자 신청 시작"].dday == "D-90"
    assert tasks["사례비 지급과 원천징수 신고"].dday == "D+7"


# ---------------------------------------------------------------- 예산
def test_budget_separates_the_gross_from_what_the_speaker_gets():
    budget = build_budget(_event(speakers=[_speaker(fee_basis="net")]))[0]
    labels = [line.label for line in budget.lines]
    assert "사례비 (지급 총액)" in labels
    assert any("연사 수령액" in label for label in labels)
    assert any("원천징수액" in label for label in labels)
    assert budget.to_speaker == 5_000_000
    assert budget.tax == 1_410_256


def test_budget_xlsx_is_written(tmp_path):
    event = _event()
    path = write_xlsx(event, build_budget(event), tmp_path)
    assert path.suffix == ".xlsx"
    assert path.stat().st_size > 2000

    from openpyxl import load_workbook
    sheet = load_workbook(path).active
    text = " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "세무 자문이 아닙니다" in text
    assert "Jane Doe" in text


# ---------------------------------------------------------------- 문서
def test_invitation_states_who_pays_what():
    event = _event()
    text = invitation_text(event, event.speakers[0])
    assert "LETTER OF INVITATION" in text
    assert "한국인공지능학회" in text
    assert "honorarium" in text.lower()
    assert "Travel:" in text and "Accommodation:" in text


def test_invitation_for_an_unpaid_speaker_says_so():
    event = _event(speakers=[_speaker(fee_krw=0, expenses_only=True)])
    text = invitation_text(event, event.speakers[0])
    assert "No honorarium will be paid" in text


def test_contract_does_not_gloss_over_withholding():
    event = _event()
    text = contract_text(event, event.speakers[0])
    assert "withholding tax" in text
    assert "Certificate of Residence" in text
    assert "before the payment date" in text


def test_contract_leaves_the_risky_clauses_blank():
    """녹화물 범위와 취소 조항은 사람이 정해야 한다."""
    text = contract_text(_event(), _speaker())
    assert "[scope]" in text and "[period]" in text
    assert "Visa refusal" in text


def test_every_document_says_it_is_a_draft():
    event = _event()
    for text in (invitation_text(event, event.speakers[0]),
                 contract_text(event, event.speakers[0])):
        assert NOT_LEGAL_ADVICE in text
        assert "법률 자문이 아닙니다" in text


def test_docx_carries_the_ai_metadata(tmp_path):
    from speaker_desk.letters import write_docx
    from shared.ai_label import METADATA_PREFIX

    from docx import Document
    path = write_docx("HELLO\nbody", tmp_path / "x.docx", "Test")
    props = Document(path).core_properties
    assert METADATA_PREFIX in (props.comments or "")


# ---------------------------------------------------------------- 브리프
def test_brief_leads_with_the_visa_warning(tmp_path):
    event = _event()
    speaker = event.speakers[0]
    tasks = build_checklist(event, speaker)
    budget = build_budget(event)[0]
    path = write_brief(event, speaker, tasks, budget, tmp_path,
                       today=date(2026, 9, 18))
    text = path.read_text(encoding="utf-8")
    assert "무비자로 강연할 수 없습니다" in text
    assert "판정이 아닙니다" in text
    assert "세무 자문이 아닙니다" in text
    assert "거주자증명서" in text


# ------------------------------------------------------------------ CLI
def test_cli_visa_warns_and_returns_two(capsys):
    cli = load_product_cli("speaker-desk")
    assert cli.main(["visa", "--paid", "--days", "5", "--waiver"]) == 2
    out = capsys.readouterr().out
    assert "단기취업 (C-4)" in out
    assert "판정이 아닙니다" in out


def test_cli_tax_shows_both_sides(capsys):
    cli = load_product_cli("speaker-desk")
    assert cli.main(["tax", "5000000", "--basis", "net"]) == 0
    out = capsys.readouterr().out
    assert "6,410,256" in out
    assert "1,410,256" in out
    assert "거주자증명서" in out


def test_cli_brief_demo_writes_one_file_per_speaker(tmp_path, capsys):
    cli = load_product_cli("speaker-desk")
    code = cli.main(["brief", "--demo", "--out", str(tmp_path),
                     "--today", "2026-09-18"])
    assert code in (0, 2)
    assert len(list(tmp_path.glob("brief_*.md"))) == 2
    assert "법률·세무 자문이 아닙니다" in capsys.readouterr().out


def test_cli_letter_makes_both_documents(tmp_path):
    cli = load_product_cli("speaker-desk")
    assert cli.main(["letter", "--demo", "--out", str(tmp_path),
                     "--speaker", "Jane Doe"]) == 0
    names = {path.name for path in tmp_path.iterdir()}
    assert "invitation_Jane_Doe.txt" in names
    assert "contract_Jane_Doe.docx" in names


def test_cli_checklist_flags_overdue(tmp_path, capsys):
    cli = load_product_cli("speaker-desk")
    code = cli.main(["checklist", "--demo", "--today", "2026-11-01"])
    assert code == 2, "행사 3주 전이면 지난 필수 일정이 있다"
    assert "지난 필수 일정" in capsys.readouterr().out


# --------------------------------------------------------------- 등록
def test_registered_as_program_4():
    program = Registry().require("speaker-desk")
    assert program.number == 4
    assert program.status == "ready"


def test_it_needs_nothing_to_run():
    """계정도 인터넷도 없다. 파시기 가장 쉬운 상품이다."""
    need = Registry().require("speaker-desk").requirements
    assert need.accounts == []
    assert need.internet == "none"
    assert need.signup_free and need.starts_today


def test_manifest_disclaims_legal_and_tax_advice():
    cautions = " ".join(Registry().require("speaker-desk").requirements.cautions)
    assert "법률·세무 자문이 아닙니다" in cautions
    assert "판정하지 않습니다" in cautions
    assert "개인정보" in cautions
