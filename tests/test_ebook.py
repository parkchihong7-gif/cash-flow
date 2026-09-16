"""products/ebook-gen 테스트.

가짜 LLM(`sample_content.fake_ask`)을 주입해 실제 Claude 호출 없이
3단계(목차 → 원고 → docx) 전 과정을 검증한다.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pytest
import yaml
from pydantic import ValidationError

from ebook_gen import docx_builder
from ebook_gen.generator import (
    BASE_DIR, EbookGenerator, parse_chapter_markdown, render_chapter_markdown,
)
from ebook_gen.sample_content import fake_ask
from ebook_gen.schema import (
    CHAPTER_MAX_CHARS, CHAPTER_MIN_CHARS, CHARS_PER_PAGE, ChapterPlan,
    LENGTH_TOLERANCE, MAX_CHAPTERS, MIN_CHAPTERS, Outline, load_outline,
    plan_budget, save_outline,
)
from shared import banned_phrases

EBOOK_DIR = Path(BASE_DIR)


# ------------------------------------------------------------------ 픽스처
@pytest.fixture
def generator():
    return EbookGenerator(model="test", ask_fn=fake_ask)


@pytest.fixture
def outline(generator):
    result, _ = generator.build_outline(
        topic="테스트 주제", audience="테스트 독자", pages=30, author="홍길동"
    )
    return result


@pytest.fixture
def drafts(generator, outline):
    """전 챕터 원고. 앞 챕터 요약을 다음 챕터로 넘긴다."""
    made, previous = [], ""
    for plan in outline.chapters:
        result = generator.write_chapter(outline, plan, previous)
        assert result.ok, f"{plan.number}장 생성 실패"
        made.append(result.draft)
        previous = result.draft.summary_text
    return made


@pytest.fixture
def built(tmp_path, outline, drafts):
    path = docx_builder.build_docx(outline, drafts, tmp_path / "ebook.docx", model="test")
    return path, outline, drafts


# ---------------------------------------------------------------- 분량 계산
@pytest.mark.parametrize("pages", [20, 25, 30, 35, 40])
def test_budget_stays_within_chapter_limits(pages):
    budget = plan_budget(pages)
    assert MIN_CHAPTERS <= budget.chapter_count <= MAX_CHAPTERS
    assert CHAPTER_MIN_CHARS <= budget.chars_per_chapter <= CHAPTER_MAX_CHARS


def test_budget_hits_40_pages_exactly():
    """사양의 기본값 40쪽은 경고 없이 맞아야 한다."""
    budget = plan_budget(40)
    assert budget.warning == ""
    assert budget.achievable_pages == 40


@pytest.mark.parametrize("pages", [10, 70])
def test_budget_warns_when_pages_unreachable(pages):
    budget = plan_budget(pages)
    assert budget.warning, f"{pages}쪽은 맞출 수 없으므로 알려야 합니다"
    assert str(budget.achievable_pages) in budget.warning


# ------------------------------------------------------------ 목차 스키마
def test_outline_has_everything_the_spec_requires(outline):
    assert len(outline.title_options) == 3
    assert len(set(outline.title_options)) == 3, "제목 3안이 서로 달라야 합니다"
    assert outline.subtitle and outline.preface
    assert MIN_CHAPTERS <= len(outline.chapters) <= MAX_CHAPTERS
    for chapter in outline.chapters:
        assert 3 <= len(chapter.subheadings) <= 5
        assert chapter.key_message
        assert CHAPTER_MIN_CHARS <= chapter.target_chars <= CHAPTER_MAX_CHARS


def test_outline_title_is_first_option(outline):
    assert outline.title == outline.title_options[0]


def _plan(**overrides) -> dict:
    payload = {
        "number": 1, "title": "제목", "key_message": "메시지",
        "subheadings": ["가", "나", "다"], "target_chars": 2000,
    }
    payload.update(overrides)
    return payload


def _outline_payload(**overrides) -> dict:
    payload = {
        "topic": "주제", "audience": "독자", "pages": 30,
        "title_options": ["1안", "2안", "3안"], "subtitle": "부제", "preface": "서문",
        "chapters": [_plan(number=i) for i in range(1, 9)],
    }
    payload.update(overrides)
    return payload


@pytest.mark.parametrize(
    ("overrides", "reason"),
    [
        ({"chapters": [_plan(number=i) for i in range(1, 8)]}, "챕터 8개 미만"),
        ({"chapters": [_plan(number=i) for i in range(1, 14)]}, "챕터 12개 초과"),
        ({"title_options": ["1안", "2안"]}, "제목 3안이 아님"),
        ({"pages": 5}, "쪽수 최소 미만"),
        ({"preface": ""}, "서문 없음"),
    ],
)
def test_invalid_outline_rejected(overrides, reason):
    with pytest.raises(ValidationError):
        Outline(**_outline_payload(**overrides))


def test_chapter_numbers_must_be_continuous():
    chapters = [_plan(number=i) for i in range(1, 9)]
    chapters[3]["number"] = 99
    with pytest.raises(ValidationError, match="빠짐없이"):
        Outline(**_outline_payload(chapters=chapters))


@pytest.mark.parametrize("target", [1499, 2501])
def test_chapter_length_outside_range_rejected(target):
    with pytest.raises(ValidationError):
        ChapterPlan(**_plan(target_chars=target))


@pytest.mark.parametrize("count", [2, 6])
def test_subheading_count_enforced(count):
    with pytest.raises(ValidationError):
        ChapterPlan(**_plan(subheadings=["가"] * count))


# ------------------------------------------------- outline.yaml 왕복 (사람 편집)
def test_outline_yaml_roundtrips(tmp_path, outline):
    path = save_outline(outline, tmp_path / "outline.yaml")
    assert load_outline(path).model_dump() == outline.model_dump()


def test_saved_outline_explains_what_can_be_edited(tmp_path, outline):
    text = save_outline(outline, tmp_path / "outline.yaml").read_text(encoding="utf-8")
    assert "고쳐도 되는 것" in text
    assert "지켜야 하는 것" in text
    assert str(MIN_CHAPTERS) in text


def test_hand_edited_outline_is_validated(tmp_path, outline):
    """사람이 규칙을 어기게 고치면 2단계에서 잡아야 한다."""
    path = save_outline(outline, tmp_path / "outline.yaml")
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    raw["chapters"] = raw["chapters"][:3]  # 챕터를 3개로 줄여버림
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")

    with pytest.raises(ValidationError):
        load_outline(path)


def test_hand_edited_title_order_changes_cover_title(tmp_path, outline):
    path = save_outline(outline, tmp_path / "outline.yaml")
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    raw["title_options"] = list(reversed(raw["title_options"]))
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")

    assert load_outline(path).title == outline.title_options[2]


# ------------------------------------------------------------ 챕터 원고
def test_every_chapter_has_the_four_required_parts(drafts):
    for draft in drafts:
        assert draft.intro_case, f"{draft.plan.number}장 도입 사례 없음"
        assert draft.sections, f"{draft.plan.number}장 핵심 설명 없음"
        assert draft.checklist, f"{draft.plan.number}장 체크리스트 없음"
        assert len(draft.summary) == 3, f"{draft.plan.number}장 요약이 3줄이 아님"


def test_chapter_sections_follow_the_outline_subheadings(outline, drafts):
    for plan, draft in zip(outline.chapters, drafts):
        assert [s["subheading"] for s in draft.sections] == plan.subheadings


@pytest.mark.parametrize("index", [0, 4, 9])
def test_each_chapter_is_1500_to_2500_chars(drafts, index):
    draft = drafts[index]
    assert CHAPTER_MIN_CHARS * 0.9 <= draft.char_count <= CHAPTER_MAX_CHARS * 1.1


def test_total_length_within_20_percent_of_outline(outline, drafts):
    """완료 기준: 총 글자 수가 목차 예상 분량 ±20% 이내."""
    total = sum(draft.char_count for draft in drafts)
    low = outline.target_chars * (1 - LENGTH_TOLERANCE)
    high = outline.target_chars * (1 + LENGTH_TOLERANCE)
    assert low <= total <= high, (
        f"총 {total:,}자가 목표 {outline.target_chars:,}자의 ±20% "
        f"({low:,.0f}~{high:,.0f}자)를 벗어났습니다"
    )
    assert outline.within_tolerance(total)


def test_previous_summary_is_passed_to_next_chapter(outline):
    """앞 챕터 요약이 다음 챕터 프롬프트에 들어가야 일관성이 유지된다."""
    seen: list[str] = []

    def spy(system, user, model="test", json_mode=False, **kw):
        seen.append(user)
        return fake_ask(system, user, model, json_mode)

    generator = EbookGenerator(model="test", ask_fn=spy)
    first = generator.write_chapter(outline, outline.chapters[0], "")
    generator.write_chapter(outline, outline.chapters[1], first.draft.summary_text)

    assert "첫 챕터입니다" in seen[0]
    assert "[앞 챕터 요약]" in seen[1]
    assert first.draft.summary[0][:15] in seen[1]


def test_failed_chapter_is_retried_alone(outline):
    """한 챕터가 실패해도 그 챕터만 다시 시도한다."""
    from shared.llm import LLMError

    state = {"calls": 0}

    def flaky(system, user, model="test", json_mode=False, **kw):
        state["calls"] += 1
        if state["calls"] == 1:
            raise LLMError("일시적인 오류")
        return fake_ask(system, user, model, json_mode)

    result = EbookGenerator(model="test", ask_fn=flaky).write_chapter(
        outline, outline.chapters[0], ""
    )
    assert result.ok
    assert result.attempts == 2
    assert result.error == ""


def test_chapter_failing_every_time_is_reported(outline):
    from shared.llm import LLMError

    def always_fail(system, user, model="test", json_mode=False, **kw):
        raise LLMError("계속 실패")

    result = EbookGenerator(model="test", ask_fn=always_fail).write_chapter(
        outline, outline.chapters[0], ""
    )
    assert not result.ok
    assert "계속 실패" in result.error


def test_banned_phrase_triggers_rewrite(outline):
    state = {"calls": 0}

    def dirty_first(system, user, model="test", json_mode=False, **kw):
        state["calls"] += 1
        if state["calls"] == 1:
            data = fake_ask(system, user, model, json_mode)
            data["intro_case"] = "이 방법은 수익 보장입니다. " + data["intro_case"]
            return data
        return fake_ask(system, user, model, json_mode)

    result = EbookGenerator(model="test", ask_fn=dirty_first).write_chapter(
        outline, outline.chapters[0], ""
    )
    assert result.attempts == 2
    assert result.clean


def test_short_chapter_triggers_rewrite(outline):
    state = {"calls": 0}

    def short_first(system, user, model="test", json_mode=False, **kw):
        state["calls"] += 1
        data = fake_ask(system, user, model, json_mode)
        if state["calls"] == 1:
            data["sections"] = [{"subheading": s["subheading"], "body": "짧음"}
                                for s in data["sections"]]
        return data

    result = EbookGenerator(model="test", ask_fn=short_first).write_chapter(
        outline, outline.chapters[0], ""
    )
    assert result.attempts == 2, "분량이 크게 모자라면 다시 써야 합니다"
    assert result.clean


# --------------------------------------------------- 마크다운 왕복 (사람 편집)
def test_chapter_markdown_roundtrips(outline, drafts):
    plan, draft = outline.chapters[0], drafts[0]
    text = render_chapter_markdown(draft)
    parsed = parse_chapter_markdown(text, plan)

    assert parsed.intro_case == draft.intro_case
    assert [s["subheading"] for s in parsed.sections] == \
           [s["subheading"] for s in draft.sections]
    assert parsed.checklist == draft.checklist
    assert parsed.summary == draft.summary


def test_hand_edited_chapter_is_reflected(outline, drafts):
    """사람이 원고를 고치면 docx 에 그대로 들어가야 한다."""
    text = render_chapter_markdown(drafts[0]).replace(
        drafts[0].checklist[0], "내가 직접 고친 체크리스트 항목"
    )
    parsed = parse_chapter_markdown(text, outline.chapters[0])
    assert parsed.checklist[0] == "내가 직접 고친 체크리스트 항목"


def test_broken_chapter_structure_is_reported(outline, drafts):
    text = render_chapter_markdown(drafts[0]).replace("## 실행 체크리스트", "## 할 일")
    with pytest.raises(ValueError, match="실행 체크리스트"):
        parse_chapter_markdown(text, outline.chapters[0])


def test_chapter_markdown_carries_ai_label(drafts):
    assert "생성형 AI" in render_chapter_markdown(drafts[0])
    assert "생성형 AI" not in render_chapter_markdown(drafts[0], ai_label=False)


# ------------------------------------------------------------------ docx
def test_docx_is_created_and_reopens(built):
    path, _, _ = built
    assert path.is_file() and path.stat().st_size > 10_000
    assert len(docx.Document(str(path)).paragraphs) > 100


def test_docx_is_a_valid_zip_with_required_parts(built):
    path, _, _ = built
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert "word/document.xml" in names
        assert "word/styles.xml" in names
        assert archive.testzip() is None, "손상된 항목이 있습니다"


def test_docx_settings_pass_schema(built):
    """python-docx 기본 템플릿의 w:zoom 스키마 위반을 고쳤는지."""
    path, _, _ = built
    with zipfile.ZipFile(path) as archive:
        settings = archive.read("word/settings.xml").decode("utf-8")
    assert 'w:val="bestFit"' not in settings
    assert 'w:zoom w:percent=' in settings


def test_docx_page_format_matches_spec(built):
    path, _, _ = built
    document = docx.Document(str(path))
    section = document.sections[0]
    assert round(section.top_margin.cm, 1) == 2.5
    assert round(section.left_margin.cm, 1) == 2.5

    normal = document.styles["Normal"]
    assert normal.font.name == docx_builder.BODY_FONT
    assert normal.font.size.pt == docx_builder.BODY_SIZE_PT
    assert normal.paragraph_format.line_spacing == docx_builder.LINE_SPACING


def test_docx_sets_korean_font_for_east_asian_text(built):
    """한글은 eastAsia 속성이 있어야 지정한 글꼴로 나온다."""
    path, _, _ = built
    with zipfile.ZipFile(path) as archive:
        styles = archive.read("word/styles.xml").decode("utf-8")
    assert "w:eastAsia" in styles
    assert docx_builder.BODY_FONT in styles


def test_docx_has_toc_and_page_number_fields(built):
    path, _, _ = built
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer = next(n for n in archive.namelist() if n.startswith("word/footer"))
        footer_xml = archive.read(footer).decode("utf-8")
    assert "TOC" in document_xml, "목차 필드가 없습니다"
    assert "PAGE" in footer_xml, "페이지 번호 필드가 없습니다"


def test_docx_contains_cover_toc_preface_and_notice(built):
    path, outline, _ = built
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    assert outline.title in text
    assert outline.subtitle in text
    assert outline.author in text
    assert "목차" in text
    assert "서문" in text
    assert "AI 활용 고지" in text


def test_docx_chapter_headings_use_built_in_styles(built):
    """TOC 가 잡히려면 내장 Heading 스타일을 써야 한다."""
    path, outline, _ = built
    document = docx.Document(str(path))
    h1 = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]

    for chapter in outline.chapters:
        assert f"{chapter.number}장. {chapter.title}" in h1
    assert "목차" in h1 and "서문" in h1 and "AI 활용 고지" in h1
    assert "실행 체크리스트" in h2


def test_docx_checklists_are_tables(built):
    path, outline, drafts = built
    document = docx.Document(str(path))
    assert len(document.tables) == len(outline.chapters)

    first = document.tables[0]
    assert first.rows[0].cells[1].text == "할 일"
    assert len(first.rows) == len(drafts[0].checklist) + 1
    assert drafts[0].checklist[0] in first.rows[1].cells[1].text


def test_docx_records_ai_metadata(built):
    """완료 기준: core properties 에 AI 메타데이터."""
    path, outline, _ = built
    core = docx.Document(str(path)).core_properties
    assert "AI-generated: true" in (core.comments or "")
    assert "tool: ebook-gen" in core.comments
    assert core.title == outline.title
    assert core.author == outline.author


def test_docx_body_text_matches_manuscript(built):
    path, _, drafts = built
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    for draft in drafts:
        assert draft.summary[0] in text, f"{draft.plan.number}장 요약이 빠졌습니다"


# ------------------------------------------------------- 판매글 / 표지 지시서
def test_sales_copy_has_three_pricing_plans(generator, outline):
    data = generator.sales_copy(outline)
    assert len(data["pricing"]) == 3
    plans = " ".join(p["plan"] for p in data["pricing"])
    assert "단독" in plans and "템플릿" in plans and "피드백" in plans
    assert all(p["price"] > 0 for p in data["pricing"])


def test_sales_copy_intro_is_around_500_chars(generator, outline):
    intro = generator.sales_copy(outline)["intro"]
    assert 350 <= len("".join(intro.split())) <= 800


def test_sales_copy_states_who_it_is_not_for(generator, outline):
    data = generator.sales_copy(outline)
    assert len(data["for_whom"]) >= 3
    assert len(data["not_for_whom"]) >= 2


def test_sales_copy_has_no_banned_phrases(generator, outline):
    data = generator.sales_copy(outline)
    text = " ".join([data["intro"], data["toc_summary"], *data["for_whom"],
                     *data["not_for_whom"]])
    assert banned_phrases.check(text) == []


def test_cover_brief_includes_image_prompt_without_text(generator, outline):
    data = generator.cover_brief(outline)
    assert data["image_prompt"]
    assert "no text" in data["image_prompt"].lower()
    assert data["mood"] and data["colors"] and data["avoid"]
    assert all(c["hex"].startswith("#") for c in data["colors"])


# ------------------------------------------------------------------ 문서
def test_readme_explains_three_stages_and_edit_points():
    text = (EBOOK_DIR / "README.md").read_text(encoding="utf-8")
    for token in ["outline", "write", "build", "outline.yaml", "chapters/"]:
        assert token in text, token
    assert "사람" in text and "검수" in text
    assert banned_phrases.check(text) == []


def test_product_manuals_are_detailed_and_clean():
    for name in ("admin.md", "client.md"):
        path = EBOOK_DIR / "docs" / name
        assert path.is_file(), name
        text = path.read_text(encoding="utf-8")
        assert len(text) > 4000, f"{name} 이 너무 짧습니다"
        assert banned_phrases.check(text) == []


def test_prompts_are_split_by_stage():
    for name in ("outline.md", "chapter.md"):
        assert (EBOOK_DIR / "prompts" / name).is_file(), name
    chapter = (EBOOK_DIR / "prompts" / "chapter.md").read_text(encoding="utf-8")
    for part in ("도입 사례", "핵심 설명", "실행 체크리스트", "요약 3줄"):
        assert part in chapter
